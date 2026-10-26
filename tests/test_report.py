"""Test cho report.py (spec docs/specs/report.md, mục 9 — TC-report-01..25)."""

from __future__ import annotations

import csv
import dataclasses
import hashlib
import io
import json
import subprocess
import sys
import time
import zipfile
from datetime import date, datetime
from pathlib import Path
from typing import Any
from zoneinfo import ZoneInfo

import pytest
from openpyxl import Workbook, load_workbook

from loc_counter import report
from loc_counter.config import AppConfig, FilterConfig, config_to_dict
from loc_counter.counter import (
    FileRow,
    LanguageRow,
    UserMetrics,
    WarningKind,
    WeekRow,
    compare_files,
    compare_languages,
    count_export,
    cross_check,
    rename_map,
    total_metrics,
    user_metrics,
    weekly_metrics,
)
from loc_counter.git_repo import CommitInfo
from loc_counter.report import (
    ExportSummary,
    ReportData,
    ReportError,
    ReportOptions,
    ReportWarning,
    build_manifest,
    clean_text,
    write_csv_reports,
    write_manifest,
    write_reports,
    write_warnings_csv,
    write_xlsx,
)
from loc_counter.snapshot import make_period, start_zip_name

BKK = ZoneInfo("Asia/Bangkok")
NOTE = "Bản export end theo user dùng để kiểm tra và đếm LOC, có thể không build/chạy được."
SHEETS = ["Tong_quan", "Theo_ngon_ngu", "Theo_user", "Theo_file", "Theo_tuan", "Canh_bao"]
GEN = datetime(2026, 10, 22, 20, 52, 0, tzinfo=BKK)
DIR_ERROR = "Thư mục kết quả không tồn tại hoặc không phải thư mục"


# ---------- dữ liệu mẫu ----------
def commit(ch: str, when: datetime, subject: str) -> CommitInfo:
    return CommitInfo(ch * 40, ch * 7, when, "dev", "dev@example.com", subject)


C_START = commit("a", datetime(2025, 12, 30, 18, 20, 15, tzinfo=BKK), "Fix login bug")
C_END = commit("f", datetime(2026, 1, 21, 10, 5, 30, tzinfo=BKK), "Release 2.3")

W_BIN = ReportWarning(WarningKind.BINARY, "b.py", "Bỏ qua file binary", ("start",))
W_CROSS = ReportWarning(WarningKind.CROSS_CHECK, "", "Kiểm tra chéo (bob): lệch",
                        ("kiểm tra chéo",))
W_UN2 = ReportWarning(WarningKind.UNATTRIBUTED, "a.py", "dòng 9", ("start", "end:alice"))
W_UN1 = ReportWarning(WarningKind.UNATTRIBUTED, "a.py", "dòng 3", ("gán tác giả",))
WARNINGS = (W_BIN, W_CROSS, W_UN2, W_UN1)


def sample(**changes: Any) -> ReportData:
    """ReportData mẫu: 2 user chọn (alice, bob), xuất riêng, 4 cảnh báo."""
    main = cross_check(10, 20, [3, 7])
    data = ReportData(
        repo="https://github.com/org/repo",
        branch="main",
        period=make_period(date(2026, 1, 1), date(2026, 1, 31)),
        c_start=C_START,
        c_end=C_END,
        users=("bob", "alice"),
        all_users=False,
        config=AppConfig(),
        options=ReportOptions(selected_users=("alice", "bob"), separate_users=True),
        loc_start=10,
        loc_end=20,
        languages=(LanguageRow("CSS", 1, 5, 0, 0), LanguageRow("JavaScript", 0, 0, 1, 13),
                   LanguageRow("Python", 1, 10, 2, 13)),
        user_rows=(UserMetrics("alice", False, 2, 7, 2, 5, 2, ("app.py",)),
                   UserMetrics("bob", False, 3, 8, 1, 8, 1, ("app.py", "util.py"))),
        files=(FileRow("util.py", None, "Python", 0, 8), FileRow("app.py", None, "Python", 10, 7),
               FileRow("new.py", "old.py", "Python", 3, 4)),
        weeks=(WeekRow("2026-W02", date(2026, 1, 5), "bob", 2, 8, 1),
               WeekRow("2026-W01", date(2025, 12, 29), "alice", 1, 5, 2)),
        cross_check=main,
        warnings=WARNINGS,
        exports=(
            ExportSummary("start", "start_2026-01-01_aaaaaaa.zip", None, None, 1, 10),
            ExportSummary("end", "end_2026-01-31_multi-2users.zip", "multi-2users",
                          ("alice", "bob"), 2, 20, main),
            ExportSummary("separate", "end_2026-01-31_alice.zip", "alice", ("alice",), 1, 13,
                          cross_check(10, 13, [3])),
            ExportSummary("separate", "end_2026-01-31_bob.zip", "bob", ("bob",), 2, 16,
                          cross_check(10, 16, [7])),
        ),
    )
    return dataclasses.replace(data, **changes) if changes else data


def rows(ws: Any) -> list[tuple]:
    """Giá trị từng hàng, bỏ None ở cuối hàng."""
    out = []
    for row in ws.iter_rows(values_only=True):
        values = list(row)
        while values and values[-1] is None:
            values.pop()
        out.append(tuple(values))
    return out


def xlsx(tmp_path: Path, data: ReportData, name: str = "r.xlsx") -> Any:
    return load_workbook(write_xlsx(tmp_path / name, data))


def read_csv(path: Path) -> list[list[str]]:
    return list(csv.reader(io.StringIO(path.read_bytes().decode("utf-8-sig"), newline="")))


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


# ---------- TC-01 .. TC-11: xlsx ----------
def test_names_and_sheets(tmp_path: Path) -> None:
    """TC-report-01: tên file §12 và đúng 6 sheet theo thứ tự."""
    assert tuple(SHEETS) == report.SHEET_NAMES
    assert (report.XLSX_NAME, report.USER_CSV_NAME, report.LANGUAGE_CSV_NAME,
            report.FILE_CSV_NAME, report.WARNINGS_CSV_NAME, report.MANIFEST_NAME) == (
        "loc_report.xlsx", "loc_by_user.csv", "loc_by_language.csv", "loc_by_file.csv",
        "warnings.csv", "manifest.json")
    assert report.end_scope("alice") == "end:alice"
    assert (report.SCOPE_START, report.SCOPE_ATTRIBUTION, report.SCOPE_CROSS_CHECK,
            report.SCOPE_IDENTITY) == ("start", "gán tác giả", "kiểm tra chéo", "định danh")
    wb = xlsx(tmp_path, sample())
    assert wb.sheetnames == SHEETS


def test_overview_sheet(tmp_path: Path) -> None:
    """TC-report-02: Tong_quan — hàng Mục/Giá trị và bảng export."""
    ws = xlsx(tmp_path, sample())["Tong_quan"]
    got = rows(ws)
    expected_head = [
        ("Mục", "Giá trị"),
        ("Repo", "https://github.com/org/repo"),
        ("Nhánh", "main"),
        ("Múi giờ", "Asia/Bangkok"),
        ("Start", "2026-01-01T00:00:00+07:00"),
        ("End", "2026-01-31T23:59:59+07:00"),
        ("C_START", "a" * 40),
        ("Snapshot start", 'aaaaaaa (2025-12-30 18:20) "Fix login bug"'),
        ("C_END", "f" * 40),
        ("Snapshot end", 'fffffff (2026-01-21 10:05) "Release 2.3"'),
        ("User", "alice, bob"),
        ("Số user được tính", 2),
        ("LOC start", 10),
        ("LOC end", 20),
        ("Chênh lệch", 10),
        ("Commit", 5),
        ("Tổng code ròng", 10),
        ("LOC end kỳ vọng", 20),
        ("Lệch kiểm tra chéo (%)", 0),
        ("Ngưỡng (%)", 0.5),
        ("Kiểm tra chéo", "Đạt"),
        ("Số cảnh báo", 4),
        ("Phiên bản công cụ", "0.1.0"),
        ("Ghi chú", NOTE),
        (),
        ("File export", "Loại", "User", "Số file", "LOC", "LOC end kỳ vọng", "Lệch (%)",
         "Kiểm tra chéo"),
        ("start_2026-01-01_aaaaaaa.zip", "Start", None, 1, 10),
        ("end_2026-01-31_multi-2users.zip", "End", "alice, bob", 2, 20, 20, 0, "Đạt"),
        ("end_2026-01-31_alice.zip", "End (riêng)", "alice", 1, 13, 13, 0, "Đạt"),
        ("end_2026-01-31_bob.zip", "End (riêng)", "bob", 2, 16, 17, 6.25, "Lệch quá ngưỡng"),
    ]
    assert got == expected_head
    assert isinstance(ws["B13"].value, int)
    assert ws["A1"].font.b and ws["A26"].font.b


def test_language_sheet(tmp_path: Path) -> None:
    """TC-report-03: Theo_ngon_ngu — cột, thứ tự, định dạng số, tiêu đề in đậm, freeze."""
    ws = xlsx(tmp_path, sample())["Theo_ngon_ngu"]
    assert rows(ws) == [
        ("Ngôn ngữ", "Số file start", "LOC start", "Số file end", "LOC end", "Chênh lệch"),
        ("Python", 1, 10, 2, 13, 3),
        ("JavaScript", 0, 0, 1, 13, 13),
        ("CSS", 1, 5, 0, 0, -5),
    ]
    assert ws["C2"].number_format == "#,##0"
    assert ws["E2"].number_format == "#,##0"
    assert ws["F4"].number_format == "+#,##0;-#,##0;0"
    assert ws["A1"].font.b and ws["F1"].font.b
    assert ws.freeze_panes == "A2"
    assert ws.auto_filter.ref == "A1:F4"
    assert ws.column_dimensions["A"].width == 12          # "JavaScript" (10) + 2
    wb = xlsx(tmp_path, sample(files=(FileRow("x" * 70, None, "Python", 0, 1),)), "w.xlsx")
    assert wb["Theo_tuan"].column_dimensions["A"].width == 10   # "2026-W01" (8) → tối thiểu 10
    assert wb["Theo_file"].column_dimensions["A"].width == 60   # 70 ký tự → tối đa 60


def test_user_sheet(tmp_path: Path) -> None:
    """TC-report-04: Theo_user — sắp (-commits, account), cột Bot, dòng Tổng luôn có."""
    users = (UserMetrics("carol", False, 2, 1, 0, 1, 0, ("x.py",)),
             UserMetrics("bot[bot]", True, 1, 4, 0, 4, 0, ("app.py",)),
             UserMetrics("alice", False, 2, 7, 2, 5, 2, ("app.py",)))
    ws = xlsx(tmp_path, sample(user_rows=users))["Theo_user"]
    total = total_metrics(list(users))
    assert rows(ws) == [
        ("User", "Bot", "Commit", "Code thêm", "Code xóa", "Code ròng", "File thay đổi",
         "Dòng thêm", "Dòng xóa"),
        ("alice", "không", 2, 5, 2, 3, 1, 7, 2),
        ("carol", "không", 2, 1, 0, 1, 1, 1, 0),
        ("bot[bot]", "có", 1, 4, 0, 4, 1, 4, 0),
        ("Tổng", None, total.commits, total.added_code, total.deleted_code, total.net_code,
         total.files_changed, total.added_lines, total.deleted_lines),
    ]
    assert total.files_changed == 2
    assert ws.auto_filter.ref == "A1:I4"
    empty = xlsx(tmp_path, sample(user_rows=()), "e.xlsx")["Theo_user"]
    assert rows(empty)[1:] == [("Tổng", None, 0, 0, 0, 0, 0, 0, 0)]


def test_file_sheet(tmp_path: Path) -> None:
    """TC-report-05: Theo_file — sắp theo path, cột Đổi tên từ, chênh lệch."""
    files = (*sample().files, FileRow("gone.py", None, "Python", 6, 0))
    ws = xlsx(tmp_path, sample(files=files))["Theo_file"]
    assert rows(ws) == [
        ("Đường dẫn", "Đổi tên từ", "Ngôn ngữ", "LOC start", "LOC end", "Chênh lệch"),
        ("app.py", None, "Python", 10, 7, -3),
        ("gone.py", None, "Python", 6, 0, -6),
        ("new.py", "old.py", "Python", 3, 4, 1),
        ("util.py", None, "Python", 0, 8, 8),
    ]


def test_week_sheet(tmp_path: Path) -> None:
    """TC-report-06: Theo_tuan — sắp (week_start, account), ngày dạng chuỗi; rỗng → tiêu đề."""
    ws = xlsx(tmp_path, sample())["Theo_tuan"]
    header = ("Tuần", "Ngày bắt đầu", "User", "Commit", "Code thêm", "Code xóa", "Code ròng")
    assert rows(ws) == [
        header,
        ("2026-W01", "2025-12-29", "alice", 1, 5, 2, 3),
        ("2026-W02", "2026-01-05", "bob", 2, 8, 1, 7),
    ]
    empty = xlsx(tmp_path, sample(weeks=()), "e.xlsx")["Theo_tuan"]
    assert rows(empty) == [header]


def test_warning_sheet(tmp_path: Path) -> None:
    """TC-report-07: Canh_bao — sắp ổn định (path, kind), Loại = giá trị chuỗi, Phạm vi gộp."""
    ws = xlsx(tmp_path, sample())["Canh_bao"]
    assert rows(ws) == [
        ("Loại", "Đường dẫn", "Nội dung", "Phạm vi"),
        ("cross_check", None, "Kiểm tra chéo (bob): lệch", "kiểm tra chéo"),
        ("unattributed", "a.py", "dòng 9", "start; end:alice"),
        ("unattributed", "a.py", "dòng 3", "gán tác giả"),
        ("binary", "b.py", "Bỏ qua file binary", "start"),
    ]
    empty = xlsx(tmp_path, sample(warnings=()), "e.xlsx")["Canh_bao"]
    assert rows(empty) == [("Loại", "Đường dẫn", "Nội dung", "Phạm vi")]


def test_xlsx_no_formula(tmp_path: Path) -> None:
    """TC-report-08: ô văn bản bắt đầu = + - @ là chuỗi, không thành công thức."""
    texts = ["=1+2", "+x.py", "-x.py", "@x.py"]
    data = sample(
        user_rows=(UserMetrics("=1+2", False, 1, 0, 0, 0, 0, ()),),
        files=tuple(FileRow(p, None, "Python", 0, 0) for p in texts[1:]),
        warnings=(ReportWarning(WarningKind.BINARY, "x", '=HYPERLINK("http://a")', ("start",)),),
    )
    wb = xlsx(tmp_path, data)
    cells = [wb["Theo_user"]["A2"], wb["Theo_file"]["A2"], wb["Theo_file"]["A3"],
             wb["Theo_file"]["A4"], wb["Canh_bao"]["C2"]]
    assert sorted(c.value for c in cells) == sorted([*texts, '=HYPERLINK("http://a")'])
    assert all(c.data_type == "s" for c in cells)


def test_control_chars_and_truncation(tmp_path: Path) -> None:
    """TC-report-09: lọc ký tự điều khiển, surrogate → U+FFFD, cắt ô xlsx 32.767 ký tự."""
    assert clean_text("a\x00b\x07c\x1bd\te\nf") == "abcd\te\nf"
    assert clean_text("x\ud800y") == "x\ufffdy"
    assert clean_text("a\ufffeb\uffffc\r") == "abc\r"
    long_msg = "x" * 40000
    data = sample(
        files=(FileRow("p\ud800.py", None, "Python", 0, 1),),
        warnings=(ReportWarning(WarningKind.BINARY, "a.py", "a\x00b\x07c\x1bd\te\nf", ("start",)),
                  ReportWarning(WarningKind.BINARY, "b.py", long_msg, ("start",))),
    )
    wb = xlsx(tmp_path, data)
    assert wb["Theo_file"]["A2"].value == "p\ufffd.py"
    assert wb["Canh_bao"]["C2"].value == "abcd\te\nf"
    cut = wb["Canh_bao"]["C3"].value
    assert len(cut) == report.XLSX_MAX_CELL_CHARS == 32767
    assert cut.endswith("…") and cut[:-1] == "x" * 32766
    path = write_warnings_csv(tmp_path / "w.csv", data.warnings)
    assert read_csv(path)[2][2] == long_msg
    # \r: xlsx chuẩn hóa thành \n (A12), CSV giữ nguyên; chuỗi chỉ có \x00 → ô trống
    cr = sample(files=(FileRow("a\rb.py", "\x00", "Python", 0, 1),
                       FileRow("c\r\nd.py", None, "Python", 0, 1)))
    ws = xlsx(tmp_path, cr, "cr.xlsx")["Theo_file"]
    assert (ws["A2"].value, ws["A3"].value) == ("a\nb.py", "c\nd.py")
    assert ws["B2"].value is None
    got = read_csv(write_csv_reports(tmp_path, cr)[2])
    assert [r[0] for r in got[1:]] == ["a\rb.py", "c\r\nd.py"]


def _entry_times(path: Path) -> list[tuple]:
    with zipfile.ZipFile(path) as zf:
        return [(i.filename, i.date_time) for i in zf.infolist()]


def test_xlsx_reproducible(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    """TC-report-10: ghi 2 lần (đồng hồ khác) → trùng SHA-256; giờ entry + core.xml = C_END."""
    originals: list[bytes] = []
    real_save = Workbook.save

    def spy_save(self: Any, target: Any) -> None:
        real_save(self, target)
        originals.append(target.getvalue())

    monkeypatch.setattr(Workbook, "save", spy_save)
    first = write_xlsx(tmp_path / "1.xlsx", sample())
    time.sleep(2.1)
    second = write_xlsx(tmp_path / "2.xlsx", sample())
    assert sha256(first) == sha256(second)
    assert len(originals) == 2 and originals[0] != originals[1]   # bản openpyxl gốc khác nhau
    stamp = C_END.committer_date.astimezone(BKK).replace(tzinfo=None).timetuple()[:6]
    names = [n for n, _ in _entry_times(first)]
    with zipfile.ZipFile(io.BytesIO(originals[0])) as zf:
        assert names == [i.filename for i in zf.infolist()]
    assert all(t == stamp for _, t in _entry_times(first))
    with zipfile.ZipFile(first) as zf:
        for info in zf.infolist():
            assert info.compress_type == zipfile.ZIP_DEFLATED
            assert info.create_system == 3
            assert info.external_attr == 0o100644 << 16
    with zipfile.ZipFile(first) as zf:
        core = zf.read("docProps/core.xml").decode("utf-8")
    utc = "2026-01-21T03:05:30Z"
    assert f">{utc}</dcterms:created>" in core
    assert f">{utc}</dcterms:modified>" in core
    assert load_workbook(first).sheetnames == SHEETS


def test_xlsx_before_1980(tmp_path: Path) -> None:
    """TC-report-11: C_END trước 1980 → giờ entry kẹp 1980-01-01, core.xml giữ ngày thật."""
    old = commit("e", datetime(1975, 1, 1, 12, 0, 0, tzinfo=BKK), "old")
    path = write_xlsx(tmp_path / "o.xlsx", sample(c_end=old))
    assert all(t == (1980, 1, 1, 0, 0, 0) for _, t in _entry_times(path))
    with zipfile.ZipFile(path) as zf:
        core = zf.read("docProps/core.xml").decode("utf-8")
    assert ">1975-01-01T05:00:00Z</dcterms:created>" in core
    assert load_workbook(path).sheetnames == SHEETS


# ---------- TC-12 .. TC-14: CSV ----------
def test_csv_format(tmp_path: Path) -> None:
    """TC-report-12: 3 CSV — BOM, \\r\\n, tiêu đề tiếng Việt, số thô, không có dòng Tổng."""
    data = sample(files=(FileRow('a,"b".py', None, "Python", 1500, 1497),))
    paths = write_csv_reports(tmp_path, data)
    assert [p.name for p in paths] == ["loc_by_user.csv", "loc_by_language.csv",
                                       "loc_by_file.csv"]
    for p in paths:
        raw = p.read_bytes()
        assert raw.startswith(b"\xef\xbb\xbf")
        assert raw.endswith(b"\r\n")
        assert raw.count(b"\n") == raw.count(b"\r\n")
    user, lang, file = (read_csv(p) for p in paths)
    assert user == [
        ["User", "Bot", "Commit", "Code thêm", "Code xóa", "Code ròng", "File thay đổi",
         "Dòng thêm", "Dòng xóa"],
        ["bob", "không", "3", "8", "1", "7", "2", "8", "1"],
        ["alice", "không", "2", "5", "2", "3", "1", "7", "2"],
    ]
    assert lang[0] == ["Ngôn ngữ", "Số file start", "LOC start", "Số file end", "LOC end",
                       "Chênh lệch"]
    assert lang[3] == ["CSS", "1", "5", "0", "0", "-5"]
    assert file == [
        ["Đường dẫn", "Đổi tên từ", "Ngôn ngữ", "LOC start", "LOC end", "Chênh lệch"],
        ['a,"b".py', "", "Python", "1500", "1497", "-3"],
    ]
    assert b'"a,""b"".py"' in paths[2].read_bytes()


def test_csv_formula_prefix(tmp_path: Path) -> None:
    """TC-report-13: CSV thêm ' trước ô văn bản bắt đầu = + - @ \\t \\r; số âm không."""
    names = ["=a", "+a", "-a", "@a", "\ta", "\ra", "a=b.py"]
    data = sample(files=tuple(FileRow(n, None, "Python", 0, 1 if n == "a=b.py" else 0)
                              for n in names) + (FileRow("z.py", None, "Python", 5, 2),))
    got = read_csv(write_csv_reports(tmp_path, data)[2])
    first = {r[0] for r in got[1:]}
    assert first == {"'=a", "'+a", "'-a", "'@a", "'\ta", "'\ra", "a=b.py", "z.py"}
    assert got[-1] == ["z.py", "", "Python", "5", "2", "-3"]


def test_warnings_csv(tmp_path: Path) -> None:
    """TC-report-14: warnings.csv — BOM, \\r\\n, cột, thứ tự như TC-07; rỗng → tiêu đề."""
    path = write_warnings_csv(tmp_path / "warnings.csv", WARNINGS)
    raw = path.read_bytes()
    assert raw.startswith(b"\xef\xbb\xbf") and raw.endswith(b"\r\n")
    assert read_csv(path) == [
        ["Loại", "Đường dẫn", "Nội dung", "Phạm vi"],
        ["cross_check", "", "Kiểm tra chéo (bob): lệch", "kiểm tra chéo"],
        ["unattributed", "a.py", "dòng 9", "start; end:alice"],
        ["unattributed", "a.py", "dòng 3", "gán tác giả"],
        ["binary", "b.py", "Bỏ qua file binary", "start"],
    ]
    empty = write_warnings_csv(tmp_path / "e.csv", ())
    assert empty.read_bytes() == "\ufeffLoại,Đường dẫn,Nội dung,Phạm vi\r\n".encode()


# ---------- TC-15 .. TC-18: manifest ----------
CC_KEYS = ["expected_loc_end", "diff_percent", "loc_end", "sum_net_code", "threshold_percent",
           "ok"]


def test_build_manifest() -> None:
    """TC-report-15: thứ tự key §12 + options, exports; giá trị từng key."""
    data = sample()
    m = build_manifest(data, GEN)
    assert list(m) == ["tool_version", "generated_at", "repo", "branch", "timezone", "start",
                       "end", "c_start", "c_end", "users", "all_users", "config", "totals",
                       "cross_check", "options", "exports"]
    assert m["tool_version"] == "0.1.0"
    assert m["generated_at"] == "2026-10-22T20:52:00+07:00"
    assert (m["repo"], m["branch"], m["timezone"]) == ("https://github.com/org/repo", "main",
                                                       "Asia/Bangkok")
    assert (m["start"], m["end"]) == ("2026-01-01T00:00:00+07:00", "2026-01-31T23:59:59+07:00")
    assert (m["c_start"], m["c_end"]) == ("a" * 40, "f" * 40)
    assert m["users"] == ["alice", "bob"]
    assert m["all_users"] is False
    assert m["config"] == config_to_dict(AppConfig())
    assert m["totals"] == {"loc_start": 10, "loc_end": 20, "loc_delta": 10}
    assert list(m["totals"]) == ["loc_start", "loc_end", "loc_delta"]
    assert list(m["cross_check"]) == CC_KEYS
    assert m["cross_check"] == {"expected_loc_end": 20, "diff_percent": 0.0, "loc_end": 20,
                                "sum_net_code": 10, "threshold_percent": 0.5, "ok": True}
    assert m["options"] == {"selected_users": ["alice", "bob"], "exclude_bots": False,
                            "separate_users": True}
    json.dumps(m)


def test_manifest_exports() -> None:
    """TC-report-16: exports[] — start không có kiểm tra chéo; zip riêng có kiểm tra riêng."""
    ex = build_manifest(sample(), GEN)["exports"]
    assert [e["file"] for e in ex] == ["start_2026-01-01_aaaaaaa.zip",
                                       "end_2026-01-31_multi-2users.zip",
                                       "end_2026-01-31_alice.zip", "end_2026-01-31_bob.zip"]
    assert ex[0] == {"role": "start", "file": "start_2026-01-01_aaaaaaa.zip", "tag": None,
                     "accounts": None, "files": 1, "loc": 10, "cross_check": None}
    assert list(ex[1]) == ["role", "file", "tag", "accounts", "files", "loc", "cross_check"]
    assert ex[1]["cross_check"] == build_manifest(sample(), GEN)["cross_check"]
    assert (ex[1]["role"], ex[1]["tag"], ex[1]["accounts"]) == ("end", "multi-2users",
                                                                ["alice", "bob"])
    assert (ex[2]["role"], ex[2]["accounts"], ex[2]["loc"]) == ("separate", ["alice"], 13)
    assert ex[2]["cross_check"]["ok"] is True
    bob = ex[3]["cross_check"]
    assert list(bob) == CC_KEYS
    assert (bob["ok"], bob["expected_loc_end"], bob["loc_end"], bob["diff_percent"]) == (
        False, 17, 16, 6.25)


def test_write_manifest(tmp_path: Path) -> None:
    """TC-report-17: JSON indent 2, ensure_ascii=False, UTF-8 không BOM, LF, newline cuối."""
    m = build_manifest(sample(users=("bé@ví_dụ",)), GEN)
    path = write_manifest(tmp_path / "manifest.json", m)
    raw = path.read_bytes()
    assert not raw.startswith(b"\xef\xbb\xbf")
    assert b"\r" not in raw and raw.endswith(b"}\n")
    text = raw.decode("utf-8")
    assert "bé@ví_dụ" in text and "\\u00e9" not in text
    assert '\n  "tool_version": "0.1.0",' in text
    assert json.loads(text) == m
    # surrogate lẻ (vd. tên nhánh không phải UTF-8 trên POSIX) → U+FFFD, không lỗi (A13)
    out = tmp_path / "out"
    out.mkdir()
    bad = "b" + chr(0xDCFF)
    written = json.loads(write_reports(out, sample(branch=bad), GEN)[-1].read_bytes())
    assert written["branch"] == "b\N{REPLACEMENT CHARACTER}"
    safe = report._json_safe({bad: [bad, (bad,), 1]})
    assert safe == {"b\N{REPLACEMENT CHARACTER}": ["b\N{REPLACEMENT CHARACTER}",
                                                    ["b\N{REPLACEMENT CHARACTER}"], 1]}


def test_generated_at_default_and_naive() -> None:
    """TC-report-18: None → giờ hiện tại có offset +07:00, không micro giây; naive → ValueError."""
    value = build_manifest(sample())["generated_at"]
    parsed = datetime.fromisoformat(value)
    assert value.endswith("+07:00") and parsed.microsecond == 0 and "." not in value
    assert abs((datetime.now(BKK) - parsed).total_seconds()) < 60
    with pytest.raises(ValueError):
        build_manifest(sample(), datetime(2026, 1, 1))


# ---------- TC-19 .. TC-22: write_reports, lỗi, nhất quán ----------
def test_write_reports_and_reproducible(tmp_path: Path, monkeypatch: pytest.MonkeyPatch
                                        ) -> None:
    """TC-report-19: 6 file đúng tên, manifest ghi cuối; 2 lần → trùng SHA-256 trừ manifest."""
    order: list[str] = []
    for name in ("write_xlsx", "write_csv_reports", "write_warnings_csv", "write_manifest"):
        real = getattr(report, name)

        def spy(*args: Any, _real: Any = real, _name: str = name) -> Any:
            order.append(_name)
            return _real(*args)

        monkeypatch.setattr(report, name, spy)
    (tmp_path / "one").mkdir()
    (tmp_path / "two").mkdir()
    first = write_reports(tmp_path / "one", sample(), GEN)
    assert order == ["write_xlsx", "write_csv_reports", "write_warnings_csv", "write_manifest"]
    second = write_reports(tmp_path / "two", sample(), datetime(2027, 3, 4, 5, 6, 7, tzinfo=BKK))
    names = ["loc_report.xlsx", "loc_by_user.csv", "loc_by_language.csv", "loc_by_file.csv",
             "warnings.csv", "manifest.json"]
    assert [p.name for p in first] == names
    assert all(p.parent == tmp_path / "one" for p in first)
    assert sorted(p.name for p in (tmp_path / "one").iterdir()) == sorted(names)
    for a, b in zip(first[:-1], second[:-1], strict=True):
        assert sha256(a) == sha256(b), a.name
    m1, m2 = (json.loads(p[-1].read_text(encoding="utf-8")) for p in (first, second))
    assert m1["generated_at"] != m2["generated_at"]
    m1.pop("generated_at")
    m2.pop("generated_at")
    assert m1 == m2


def test_write_errors(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    """TC-report-20: thư mục không tồn tại / là file; lỗi ghi → ReportError tiếng Việt."""
    missing = tmp_path / "missing"
    with pytest.raises(ReportError, match=DIR_ERROR) as info:
        write_reports(missing, sample(), GEN)
    assert str(missing) in str(info.value)
    afile = tmp_path / "afile"
    afile.write_text("x")
    with pytest.raises(ReportError, match=DIR_ERROR):
        write_reports(afile, sample(), GEN)
    out = tmp_path / "out"
    out.mkdir()
    with pytest.raises(ValueError):
        write_reports(out, sample(), datetime(2026, 1, 1))
    assert list(out.iterdir()) == []

    def deny(self: Path, data: bytes) -> int:
        raise PermissionError(13, "Permission denied")

    monkeypatch.setattr(Path, "write_bytes", deny)
    with pytest.raises(ReportError, match="Không ghi được file báo cáo") as info:
        write_reports(out, sample(), GEN)
    assert "loc_report.xlsx" in str(info.value)
    assert isinstance(info.value.__cause__, PermissionError)
    for call in (lambda: write_csv_reports(out, sample()),
                 lambda: write_warnings_csv(out / "w.csv", WARNINGS),
                 lambda: write_manifest(out / "m.json", {"a": 1})):
        with pytest.raises(ReportError, match="Không ghi được file báo cáo"):
            call()


def test_report_data_consistency() -> None:
    """TC-report-21: cross_check không khớp loc_start / loc_end → ValueError."""
    with pytest.raises(ValueError):
        sample(loc_start=11)
    with pytest.raises(ValueError):
        sample(loc_end=21)
    assert sample().loc_delta == 10


def _keys(value: Any) -> list[str]:
    if isinstance(value, dict):
        return [name for key, v in value.items() for name in (key, *_keys(v))]
    if isinstance(value, list):
        return [name for v in value for name in _keys(v)]
    return []


def test_no_secret_fields() -> None:
    """TC-report-22: không có trường / key manifest nào mang tên bí mật."""
    names = [f.name for cls in (ReportData, ReportOptions, ExportSummary, ReportWarning)
             for f in dataclasses.fields(cls)]
    names += _keys(build_manifest(sample(), GEN))
    bad = ("token", "password", "secret", "auth")
    assert names and not [n for n in names if any(b in n.lower() for b in bad)]


# ---------- TC-23: §14 tích hợp ----------
def test_section14_reports(section14_repo, cache_root: Path, tmp_path: Path) -> None:
    """TC-report-23: §14 (all) — xlsx, manifest đúng số liệu; chạy 2 lần trùng SHA-256."""
    from test_attribution import A, B
    from test_counter_metrics import s14_ctx

    s = s14_ctx(section14_repo, cache_root)
    filters = FilterConfig()
    start = count_export(s.exports["start"], s.repo.read_blobs, filters)
    end = count_export(s.exports["all"], s.repo.read_blobs, filters)
    metrics = user_metrics(s.attr, s.summaries, [A, B])
    check = cross_check(start.total_code, end.total_code, [m.net_code for m in metrics])
    period = make_period(date(2026, 1, 1), date(2026, 1, 31))
    data = ReportData(
        repo=s.repo.source, branch="main", period=period, c_start=s.snaps.c_start,
        c_end=s.snaps.c_end, users=(A, B), all_users=True, config=AppConfig(),
        options=ReportOptions(), loc_start=start.total_code, loc_end=end.total_code,
        languages=tuple(compare_languages(start, end)), user_rows=tuple(metrics),
        files=tuple(compare_files(start, end, rename_map(s.attr, [f.path for f in end.files]))),
        weeks=tuple(weekly_metrics(s.attr, s.commits, [A, B], "Asia/Bangkok")),
        cross_check=check,
        exports=(ExportSummary("start", start_zip_name(period, s.snaps.c_start), None, None,
                               len(s.exports["start"].files), start.total_code),
                 ExportSummary("end", "end_2026-01-31_all.zip", "all", None,
                               len(s.exports["all"].files), end.total_code, check)),
    )
    (tmp_path / "r1").mkdir()
    (tmp_path / "r2").mkdir()
    first = write_reports(tmp_path / "r1", data, GEN)
    second = write_reports(tmp_path / "r2", data, GEN)
    for a, b in zip(first, second, strict=True):
        assert sha256(a) == sha256(b), a.name
    wb = load_workbook(first[0])
    assert wb.sheetnames == SHEETS
    overview = dict(r[:2] for r in rows(wb["Tong_quan"])[1:24] if r)
    assert (overview["LOC start"], overview["LOC end"], overview["Chênh lệch"],
            overview["Commit"]) == (10, 20, 10, 4)
    assert overview["User"] == "Tất cả user" and "Ghi chú" not in overview
    users = {r[0]: r[3:6] for r in rows(wb["Theo_user"])[1:]}
    assert users == {A: (5, 2, 3), B: (8, 1, 7), "Tổng": (13, 3, 10)}
    weeks = rows(wb["Theo_tuan"])[1:]
    assert [w[0] for w in weeks] == ["2026-W02", "2026-W02"]
    manifest = json.loads(first[-1].read_text(encoding="utf-8"))
    assert manifest["totals"] == {"loc_start": 10, "loc_end": 20, "loc_delta": 10}
    assert manifest["cross_check"]["diff_percent"] == 0.0
    assert manifest["cross_check"]["ok"] is True
    assert manifest["c_start"] == s.snaps.c_start.sha
    assert manifest["users"] == sorted([A, B]) and manifest["all_users"] is True


# ---------- TC-24, TC-25 ----------
def test_no_git_imports() -> None:
    """TC-report-24: import loc_counter.report không kéo theo module git / pipeline."""
    code = ("import sys, loc_counter.report; "
            "print(' '.join(m for m in sys.modules if m.startswith('loc_counter')))")
    out = subprocess.run([sys.executable, "-c", code], capture_output=True, text=True,
                         check=True).stdout.split()
    assert "loc_counter.report" in out
    for name in ("git_repo", "snapshot", "attribution", "builder", "identity", "pipeline"):
        assert f"loc_counter.{name}" not in out


def test_empty_start_all_users(tmp_path: Path) -> None:
    """TC-report-25: C_START rỗng, không commit, tất cả user → không Ghi chú; loại trừ bot."""
    check = cross_check(0, 20, [])
    data = sample(c_start=None, loc_start=0, user_rows=(), weeks=(), all_users=True,
                  options=ReportOptions(), cross_check=check,
                  exports=(ExportSummary("start", "start_2026-01-01_empty.zip", None, None, 0, 0),
                           ExportSummary("end", "end_2026-01-31_all.zip", "all", None, 3, 20,
                                         check)))
    got = dict(r[:2] for r in rows(xlsx(tmp_path, data)["Tong_quan"])[1:23] if r)
    assert got["C_START"] == "(rỗng)"
    assert got["Snapshot start"] == "Không có commit trước start (LOC start = 0)"
    assert got["User"] == "Tất cả user"
    assert got["Commit"] == 0
    assert "Ghi chú" not in got
    end_row = rows(xlsx(tmp_path, data, "a.xlsx")["Tong_quan"])[-1]
    assert end_row[:3] == ("end_2026-01-31_all.zip", "End", "Tất cả user")
    bots = sample(c_start=None, loc_start=0, all_users=True, cross_check=check,
                  options=ReportOptions(exclude_bots=True),
                  exports=(ExportSummary("end", "end_2026-01-31_all.zip", "all", ("alice",), 3,
                                         20, check),))
    sheet = rows(xlsx(tmp_path, bots, "b.xlsx")["Tong_quan"])
    got = dict(r[:2] for r in sheet[1:25] if r)
    assert got["User"] == "Tất cả user (loại trừ bot)"
    assert sheet[-1][:3] == ("end_2026-01-31_all.zip", "End", "Tất cả user (loại trừ bot)")
    assert got["Ghi chú"] == NOTE
    assert build_manifest(data, GEN)["c_start"] is None
