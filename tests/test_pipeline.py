"""Test cho pipeline.py (spec docs/specs/pipeline.md, mục 9)."""

from __future__ import annotations

import base64
import contextlib
import dataclasses
import hashlib
import json
import logging
import os
import shutil
import sys
import threading
import zipfile
from collections.abc import Callable
from datetime import date, datetime, time, timedelta, timezone
from pathlib import Path
from typing import Any

import pytest
from conftest import (
    Section14,
    _build_section14,
    build_cross_check_mismatch,
    build_systematic_mismatch,
)
from openpyxl import load_workbook
from test_identity import FakeResponse, FakeSession

from loc_counter import pipeline
from loc_counter.builder import user_tag
from loc_counter.config import AppConfig, FilterConfig
from loc_counter.counter import WarningKind, count_directory
from loc_counter.git_repo import GitError, GitRepo
from loc_counter.identity import (
    AccountResolver,
    GitHubLoginResolver,
    IdentityMap,
    IdentityRule,
    UserSummary,
)
from loc_counter.pipeline import (
    STEP_LABELS,
    AnalysisRequest,
    AnalysisResult,
    PeriodContext,
    PipelineError,
    Step,
    make_resolver,
    match_users,
    prepare_period,
    result_dir_name,
    run_analysis,
    run_log,
)
from loc_counter.report import ReportError
from loc_counter.snapshot import Period, SnapshotError, make_period

TOKEN = "ghp_SECRET123"
TOKEN_B64 = base64.b64encode(f"x-access-token:{TOKEN}".encode()).decode("ascii")
NL = chr(10)
JAN = (date(2026, 1, 1), date(2026, 1, 31))
BKK = timezone(timedelta(hours=7))
NOW = datetime(2026, 10, 24, 9, 0, 0, tzinfo=BKK)
IDS = IdentityMap((IdentityRule("alice", email="alice@example.com"),
                   IdentityRule("bob", email="bob@example.com")))
BOT = "dependabot[bot]"
ALICE = ("alice", "alice@example.com")
FULL = FilterConfig(export_mode="full")
REPORT_NAMES = ["loc_report.xlsx", "loc_by_user.csv", "loc_by_language.csv",
                "loc_by_file.csv", "loc_by_pr.csv", "warnings.csv", "manifest.json"]
ALL_DIR = "repo_main_2026-01-01_2026-01-31_all"


# ---------- helper ----------
def jan() -> Period:
    return make_period(*JAN)


def open_repo(src: Any, cache_root: Path) -> GitRepo:
    return GitRepo.open(str(src.path), cache_root=cache_root)


def analyze(repo: GitRepo, out_root: Path, users: tuple[str, ...] | None = None, *,
            period: Period | None = None, config: AppConfig | None = None,
            exclude_bots: bool = False, separate: bool = False,
            resolver: AccountResolver | None = None,
            progress: pipeline.ProgressCallback | None = None,
            context: PeriodContext | None = None, now: datetime = NOW,
            secrets: tuple[str | None, ...] = ()) -> AnalysisResult:
    """Cách gọi chuẩn: run_log bọc run_analysis."""
    request = AnalysisRequest("main", period or jan(), config or AppConfig(), out_root, users,
                              exclude_bots, separate)
    with run_log(out_root, secrets=secrets):
        return run_analysis(repo, request, resolver or AccountResolver(IDS), progress,
                            context=context, now=lambda: now)


def leftovers(out_root: Path) -> list[str]:
    return sorted(p.name for p in out_root.iterdir()
                  if p.name.startswith((pipeline.TMP_PREFIX, pipeline.OLD_PREFIX)))


def zip_files(path: Path) -> dict[str, bytes]:
    with zipfile.ZipFile(path) as zf:
        return {name: zf.read(name) for name in zf.namelist()}


def rows_of(result: AnalysisResult) -> dict[str, Any]:
    return {r.account: r for r in result.data.user_rows}


def manifest_of(out_dir: Path) -> dict[str, Any]:
    return json.loads((out_dir / "manifest.json").read_text(encoding="utf-8"))


def digests(out_dir: Path) -> dict[str, str]:
    return {p.name: hashlib.sha256(p.read_bytes()).hexdigest() for p in sorted(out_dir.iterdir())}


def warnings_of(result: AnalysisResult, kind: WarningKind) -> list[Any]:
    return [w for w in result.data.warnings if w.kind == kind]


class Spy:
    """Bọc một hàm của module pipeline, ghi lại tham số các lần gọi."""

    def __init__(self, monkeypatch: pytest.MonkeyPatch, owner: Any, name: str) -> None:
        self.calls: list[tuple[tuple[Any, ...], dict[str, Any]]] = []
        self._orig = getattr(owner, name)
        monkeypatch.setattr(owner, name, self)

    def __call__(self, *args: Any, **kwargs: Any) -> Any:
        self.calls.append((args, kwargs))
        return self._orig(*args, **kwargs)


def forbid(monkeypatch: pytest.MonkeyPatch, *names: str) -> None:
    """Các hàm nặng của pipeline không được gọi (kiểm tra trước việc nặng)."""
    for name in names:
        def fail(*_a: Any, _name: str = name, **_k: Any) -> Any:
            raise AssertionError(f"{_name} không được gọi")
        monkeypatch.setattr(pipeline, name, fail)


@pytest.fixture
def s14(section14_repo, cache_root: Path) -> Callable[..., tuple[Section14, GitRepo]]:
    """Repo §14 tên `repo` (thư mục kết quả `repo_…`) và GitRepo đã mở."""

    def factory(bot: bool = False) -> tuple[Section14, GitRepo]:
        sec = section14_repo(bot=bot, name="repo")
        return sec, open_repo(sec.src, cache_root)

    return factory


def github_repo(local: GitRepo) -> GitRepo:
    """GitRepo trỏ bare cache có sẵn nhưng nguồn là GitHub https + token (không mạng)."""
    return GitRepo(local.git_dir, "https://github.com/org/repo", TOKEN)


# ---------- TC-01 .. TC-06: §14 ----------
def test_all_users(s14, tmp_path: Path) -> None:
    """TC-pipeline-01: §14 tất cả user — số liệu, kiểm tra chéo, thư mục, export."""
    _, repo = s14()
    r = analyze(repo, tmp_path / "out")
    d = r.data
    assert (d.loc_start, d.loc_end) == (10, 20)
    rows = rows_of(r)
    assert (rows["alice"].added_code, rows["alice"].deleted_code, rows["alice"].net_code) == (
        5, 2, 3)
    assert (rows["bob"].added_code, rows["bob"].deleted_code, rows["bob"].net_code) == (8, 1, 7)
    assert d.cross_check.diff_percent == 0 and d.cross_check.ok
    assert not warnings_of(r, WarningKind.CROSS_CHECK)
    assert r.out_dir == (tmp_path / "out" / ALL_DIR).resolve()
    assert [e.role for e in d.exports] == ["start", "end"]
    assert d.exports[1].accounts is None and d.exports[1].tag == "all"
    assert d.all_users is True and d.users == ("alice", "bob")
    assert d.options.selected_users is None


def test_only_alice(s14, tmp_path: Path) -> None:
    """TC-pipeline-02: chỉ alice — 10/13, zip không có util.py, thư mục …_alice."""
    _, repo = s14()
    r = analyze(repo, tmp_path / "out", ("alice",))
    assert (r.data.loc_start, r.data.loc_end) == (10, 13)
    zip_path = r.out_dir / "end_2026-01-31_alice.zip"
    assert zip_path.is_file() and "util.py" not in zip_files(zip_path)
    assert r.out_dir.name == "repo_main_2026-01-01_2026-01-31_alice"
    assert [u.account for u in r.data.user_rows] == ["alice"]
    assert r.data.all_users is False and r.data.options.selected_users == ("alice",)


def test_only_bob(s14, tmp_path: Path) -> None:
    """TC-pipeline-03: chỉ bob — 10/17, app.py không có thay đổi của alice."""
    _, repo = s14()
    r = analyze(repo, tmp_path / "out", ("bob",))
    assert (r.data.loc_start, r.data.loc_end) == (10, 17)
    app = zip_files(r.end_zips[0])["app.py"].decode("utf-8")
    assert "a1 = 1" not in app and "v0 = 0" in app and "v9 = 9" not in app
    assert "util.py" in zip_files(r.end_zips[0])


def test_alice_and_bob(s14, tmp_path: Path) -> None:
    """TC-pipeline-04: alice + bob — 10/20, tag multi-2users, nội dung trùng C_END."""
    sec, repo = s14()
    r = analyze(repo, tmp_path / "out", ("bob", "alice"))
    assert (r.data.loc_start, r.data.loc_end) == (10, 20)
    assert r.data.exports[1].tag == "multi-2users"
    assert r.out_dir.name == "repo_main_2026-01-01_2026-01-31_multi-2users"
    files = zip_files(r.out_dir / "end_2026-01-31_multi-2users.zip")
    assert files["app.py"] == sec.app_end.encode("utf-8")
    assert files["util.py"] == sec.util.encode("utf-8")
    assert r.data.options.selected_users == ("alice", "bob")
    assert r.data.users == ("alice", "bob")


def test_separate_users(s14, tmp_path: Path) -> None:
    """TC-pipeline-05: xuất riêng — 3 zip end, số liệu chính theo bản gộp; 1 user → 1 zip."""
    _, repo = s14()
    r = analyze(repo, tmp_path / "out", ("alice", "bob"), separate=True)
    ex = r.data.exports
    assert [e.role for e in ex] == ["start", "end", "separate", "separate"]
    assert [e.tag for e in ex[1:]] == ["multi-2users", "alice", "bob"]
    assert [p.name for p in r.end_zips] == [
        "end_2026-01-31_multi-2users.zip", "end_2026-01-31_alice.zip", "end_2026-01-31_bob.zip"]
    assert (r.data.loc_start, r.data.loc_end) == (10, 20)
    assert [e.loc for e in ex[2:]] == [13, 17]
    assert all(e.cross_check is not None and e.cross_check.ok for e in ex[1:])
    assert [p.name for p in r.end_zips] == [e.file_name for e in ex[1:]]

    one = analyze(repo, tmp_path / "out1", ("alice",), separate=True)
    assert len(one.end_zips) == 1 and [e.role for e in one.data.exports] == ["start", "end"]


def test_all_users_with_bot(s14, tmp_path: Path) -> None:
    """TC-pipeline-06: "Tất cả" + bot — có/không loại bot; repo không bot + exclude_bots."""
    _, repo = s14(bot=True)
    bot_tag = user_tag([BOT], taken=["all", "alice", "bob"])
    a = analyze(repo, tmp_path / "a", separate=True)
    assert BOT in a.data.users
    assert f"end_2026-01-31_{bot_tag}.zip" in [p.name for p in a.end_zips]
    assert a.data.exports[1].accounts is None

    b = analyze(repo, tmp_path / "b", exclude_bots=True, separate=True)
    assert all(BOT not in (e.accounts or ()) for e in b.data.exports)
    assert [e.tag for e in b.data.exports[1:]] == ["all", "alice", "bob"]
    assert b.data.exports[1].accounts == ("alice", "bob")
    assert BOT not in rows_of(b) and b.data.users == ("alice", "bob")
    assert b.data.cross_check.ok


def test_all_users_exclude_bots_without_bot(s14, tmp_path: Path) -> None:
    """TC-pipeline-06: (c) repo không có bot + exclude_bots → end = C_END nguyên bản (D21)."""
    _, repo = s14()
    c = analyze(repo, tmp_path / "c", exclude_bots=True)
    assert c.data.exports[1].accounts is None and c.data.exports[1].tag == "all"
    assert (c.data.loc_start, c.data.loc_end) == (10, 20)


# ---------- TC-07 .. TC-13: thư mục kết quả ----------
def test_output_files(s14, tmp_path: Path) -> None:
    """TC-pipeline-07: đủ file §12, run.log ngoài thư mục kết quả, không sót thư mục tạm."""
    sec, repo = s14()
    out = tmp_path / "out"
    r = analyze(repo, out)
    sha7 = r.snapshots.c_start.short_sha
    names = sorted(p.name for p in r.out_dir.iterdir())
    assert names == sorted([f"start_2026-01-01_{sha7}.zip", "end_2026-01-31_all.zip",
                            *REPORT_NAMES])
    wb = load_workbook(r.out_dir / "loc_report.xlsx")
    assert wb.sheetnames == ["Tong_quan", "Theo_ngon_ngu", "Theo_user", "Theo_file",
                             "Theo_tuan", "Theo_PR", "Canh_bao"]
    for name in ("loc_by_user.csv", "loc_by_language.csv", "loc_by_file.csv", "loc_by_pr.csv"):
        assert (r.out_dir / name).read_bytes().startswith(b"\xef\xbb\xbf")
    m = manifest_of(r.out_dir)
    for key in ("tool_version", "generated_at", "repo", "branch", "timezone", "start", "end",
                "c_start", "c_end", "users", "all_users", "config", "totals", "cross_check"):
        assert key in m
    assert m["generated_at"] == NOW.isoformat()
    assert (out / "run.log").is_file() and not (r.out_dir / "run.log").exists()
    assert leftovers(out) == []


def test_reproducible(s14, tmp_path: Path) -> None:
    """TC-pipeline-08: hai lần chạy, now khác nhau → file trùng SHA-256 trừ manifest."""
    _, repo = s14()
    r1 = analyze(repo, tmp_path / "o1", ("alice", "bob"), separate=True)
    r2 = analyze(repo, tmp_path / "o2", ("alice", "bob"), separate=True,
                 now=NOW + timedelta(hours=5))
    d1, d2 = digests(r1.out_dir), digests(r2.out_dir)
    assert d1.keys() == d2.keys()
    assert {k: v for k, v in d1.items() if k != "manifest.json"} == {
        k: v for k, v in d2.items() if k != "manifest.json"}
    m1, m2 = manifest_of(r1.out_dir), manifest_of(r2.out_dir)
    assert m1.pop("generated_at") != m2.pop("generated_at")
    assert m1 == m2


def test_rerun_replaces_directory(s14, tmp_path: Path) -> None:
    """TC-pipeline-09: chạy lại cùng out → thư mục bị thay hoàn toàn, không còn .loc-old-*."""
    _, repo = s14()
    out = tmp_path / "out"
    r1 = analyze(repo, out)
    before = digests(r1.out_dir)
    (r1.out_dir / "stray.txt").write_text("x", encoding="utf-8")
    r2 = analyze(repo, out, now=NOW + timedelta(days=1))
    assert r2.out_dir == r1.out_dir
    after = digests(r2.out_dir)
    assert "stray.txt" not in after
    assert {k: v for k, v in after.items() if k != "manifest.json"} == {
        k: v for k, v in before.items() if k != "manifest.json"}
    assert leftovers(out) == []


@pytest.mark.parametrize("variant", ["no_manifest", "broken_json", "missing_key", "file"])
def test_target_not_tool_directory(s14, tmp_path: Path, monkeypatch: pytest.MonkeyPatch,
                                   variant: str) -> None:
    """TC-pipeline-10: đích có sẵn không phải của tool → PipelineError trước việc nặng."""
    _, repo = s14()
    out = tmp_path / "out"
    out.mkdir()
    target = out / ALL_DIR
    if variant == "file":
        target.write_text("không phải thư mục", encoding="utf-8")
    else:
        target.mkdir()
        (target / "keep.txt").write_text("giữ", encoding="utf-8")
        if variant == "broken_json":
            (target / "manifest.json").write_text("{không phải json", encoding="utf-8")
        elif variant == "missing_key":
            (target / "manifest.json").write_text(
                json.dumps({"tool_version": "0.1.0", "c_end": "x"}), encoding="utf-8")
    snapshot_before = (sorted(p.name for p in target.iterdir()) if target.is_dir()
                       else target.read_text(encoding="utf-8"))
    forbid(monkeypatch, "attribute_period", "plan_export")
    with pytest.raises(PipelineError) as exc:
        analyze(repo, out)
    assert str(target.resolve()) in str(exc.value)
    if variant == "file":
        assert "file trùng tên" in str(exc.value)
        assert target.read_text(encoding="utf-8") == snapshot_before
    else:
        assert "không phải kết quả của LOC Counter" in str(exc.value)
        assert sorted(p.name for p in target.iterdir()) == snapshot_before
    assert leftovers(out) == []


@pytest.mark.skipif(os.name != "nt", reason="chỉ Windows khóa thư mục có file đang mở")
def test_old_directory_locked(s14, tmp_path: Path) -> None:
    """TC-pipeline-11: thư mục cũ có file đang mở → PipelineError, giữ thư mục cũ."""
    _, repo = s14()
    out = tmp_path / "out"
    r1 = analyze(repo, out)
    before = digests(r1.out_dir)
    with open(r1.out_dir / "warnings.csv", "rb"), pytest.raises(PipelineError) as exc:
        analyze(repo, out)
    assert "Không thay thế được thư mục kết quả cũ" in str(exc.value)
    assert digests(r1.out_dir) == before
    assert leftovers(out) == []


def test_failure_midway(s14, tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    """TC-pipeline-12: lỗi giữa chừng → ném lại đúng loại, thư mục cũ nguyên, log ERROR."""
    _, repo = s14()
    out = tmp_path / "out"
    r1 = analyze(repo, out)
    alice_dir = analyze(repo, out, ("alice",)).out_dir
    before, before_alice = digests(r1.out_dir), digests(alice_dir)

    def boom(*_a: Any, **_k: Any) -> Any:
        raise ReportError("Lỗi ghi báo cáo giả")

    with monkeypatch.context() as m:
        m.setattr(pipeline, "write_reports", boom)
        with pytest.raises(ReportError):
            analyze(repo, out)
    assert digests(r1.out_dir) == before
    assert leftovers(out) == []
    log = (out / "run.log").read_text(encoding="utf-8")
    assert " ERROR " in log and "Traceback" in log and "Lỗi ghi báo cáo giả" in log

    def interrupt(*_a: Any, **_k: Any) -> Any:
        raise KeyboardInterrupt

    monkeypatch.setattr(pipeline, "build_user_export", interrupt)
    with pytest.raises(KeyboardInterrupt):
        analyze(repo, out, ("alice",))
    assert digests(alice_dir) == before_alice
    assert leftovers(out) == []
    assert "Traceback" in (out / "run.log").read_text(encoding="utf-8")


def test_out_root(s14, tmp_path: Path) -> None:
    """TC-pipeline-13: out_root chưa có (lồng 2 cấp) → tạo; là file → PipelineError."""
    _, repo = s14()
    nested = tmp_path / "a" / "b"
    r = analyze(repo, nested)
    assert r.out_dir.parent == nested.resolve() and (nested / "run.log").is_file()

    as_file = tmp_path / "out.txt"
    as_file.write_text("x", encoding="utf-8")
    with pytest.raises(PipelineError, match="Thư mục đầu ra là một file"), run_log(as_file):
        pass
    request = AnalysisRequest("main", jan(), AppConfig(), as_file)
    with pytest.raises(PipelineError, match="Thư mục đầu ra là một file"):
        run_analysis(repo, request, AccountResolver(IDS))


# ---------- TC-14, TC-15, TC-18: kỳ đặc biệt ----------
def test_empty_start(s14, tmp_path: Path) -> None:
    """TC-pipeline-14: C_START rỗng → zip start rỗng, loc_start 0, manifest c_start null."""
    _, repo = s14()
    period = make_period(date(2025, 1, 1), JAN[1])
    r = analyze(repo, tmp_path / "out", period=period)
    assert r.start_zip.name == "start_2025-01-01_empty.zip"
    assert zip_files(r.start_zip) == {}
    assert r.data.loc_start == 0 and r.data.c_start is None
    assert manifest_of(r.out_dir)["c_start"] is None
    assert r.data.cross_check.ok
    assert "admin@example.com" in r.data.users
    assert r.out_dir.name == "repo_main_2025-01-01_2026-01-31_all"


def test_no_commit_in_period(s14, tmp_path: Path) -> None:
    """TC-pipeline-15: kỳ không có commit → users rỗng, kiểm tra chéo đạt; chọn user → lỗi."""
    _, repo = s14()
    period = make_period(date(2025, 12, 20), date(2025, 12, 31))
    r = analyze(repo, tmp_path / "out", period=period)
    assert r.data.users == () and r.data.user_rows == ()
    assert r.data.loc_end == r.data.loc_start == 10
    assert r.data.cross_check.ok
    assert load_workbook(r.out_dir / "loc_report.xlsx")["Theo_tuan"].max_row == 1
    with pytest.raises(PipelineError, match="Không có user nào có commit trong kỳ"):
        analyze(repo, tmp_path / "out2", ("alice",), period=period)


def test_no_data_before_end(s14, tmp_path: Path) -> None:
    """TC-pipeline-18: không có commit trước end → SnapshotError, out_root chỉ có run.log."""
    _, repo = s14()
    out = tmp_path / "out"
    period = make_period(date(2020, 1, 1), date(2020, 1, 31))
    with pytest.raises(SnapshotError, match="Không có dữ liệu trong khoảng thời gian"):
        analyze(repo, out, period=period)
    assert sorted(p.name for p in out.iterdir()) == ["run.log"]


# ---------- TC-16, TC-17: chọn user ----------
def summaries(*accounts: str, bots: tuple[str, ...] = ()) -> list[UserSummary]:
    return [UserSummary(a, 1, 0, [], a in bots) for a in accounts]


def test_match_users() -> None:
    """TC-pipeline-16: khớp chính xác / casefold, mơ hồ, trùng, rỗng, cắt danh sách."""
    users = summaries("alice", "Bob", "x[bot]", bots=("x[bot]",))
    assert match_users(["ALICE"], users) == ("alice",)
    assert match_users([" alice ", "alice"], users) == ("alice",)
    assert match_users(["x[bot]", "bob"], users) == ("Bob", "x[bot]")
    with pytest.raises(PipelineError) as exc:
        match_users(["zz"], users)
    assert str(exc.value) == "Không tìm thấy user 'zz' trong kỳ. User trong kỳ: alice, Bob, x[bot]"

    ambiguous = summaries("Bob", "BOB", "alice")
    with pytest.raises(PipelineError) as exc:
        match_users(["bob"], ambiguous)
    assert str(exc.value) == ("User 'bob' khớp nhiều account chỉ khác hoa thường: Bob, BOB. "
                              "Hãy nhập chính xác.")
    assert match_users(["Bob"], summaries("Bob", "bob")) == ("Bob",)

    with pytest.raises(PipelineError) as exc:
        match_users([], users)
    assert str(exc.value) == "Chưa chọn user nào."
    with pytest.raises(PipelineError) as exc:
        match_users(["alice"], [])
    assert str(exc.value) == "Không có user nào có commit trong kỳ; không thể chọn user."

    many = summaries(*(f"u{i:02d}" for i in range(26)))
    with pytest.raises(PipelineError) as exc:
        match_users(["zz"], many)
    listed = ", ".join(f"u{i:02d}" for i in range(20))
    assert str(exc.value) == (f"Không tìm thấy user 'zz' trong kỳ. User trong kỳ: {listed}"
                              ", … và 6 user khác")


def test_explicit_bot_with_exclude_bots(s14, tmp_path: Path) -> None:
    """TC-pipeline-17: bot chọn tường minh thắng exclude_bots; export dựng cho bot."""
    _, repo = s14(bot=True)
    r = analyze(repo, tmp_path / "out", (BOT,), exclude_bots=True)
    assert r.data.users == (BOT,) and r.data.exports[1].accounts == (BOT,)
    assert [u.account for u in r.data.user_rows] == [BOT]
    assert r.data.options.exclude_bots is True


# ---------- TC-19 .. TC-23: cảnh báo ----------
def test_cross_check_mismatch(make_repo, cache_root: Path, tmp_path: Path) -> None:
    """TC-pipeline-19: đổi tên .txt → .py trong kỳ → cảnh báo kiểm tra chéo, không lỗi."""
    src = build_cross_check_mismatch(make_repo("repo"))
    r = analyze(open_repo(src, cache_root), tmp_path / "out")
    assert not r.data.cross_check.ok
    cc = warnings_of(r, WarningKind.CROSS_CHECK)
    assert len(cc) == 1 and cc[0].scopes == ("kiểm tra chéo",)
    csv_text = (r.out_dir / "warnings.csv").read_text(encoding="utf-8-sig")
    line = next(row for row in csv_text.splitlines() if "cross_check" in row)
    assert "loc_start 10 + tổng code ròng +0 = 10, loc_end = 20" in line
    assert "kiểm tra chéo" in line
    assert manifest_of(r.out_dir)["cross_check"]["ok"] is False


@pytest.mark.parametrize("kind", ["vendor", "symlink"])
def test_systematic_cross_check_mismatch(make_repo, cache_root: Path, tmp_path: Path,
                                         kind: str) -> None:
    """TC-pipeline-37: chuyển vào vendor/ / thay bằng symlink → lệch kiểm tra chéo, không lỗi."""
    src = build_systematic_mismatch(make_repo("repo"), kind)
    r = analyze(open_repo(src, cache_root), tmp_path / "out")
    cc = r.data.cross_check
    assert (cc.loc_start, cc.sum_net_code, cc.loc_end) == (20, 0, 10)
    assert cc.ok is False and cc.diff_percent == 100.0
    assert rows_of(r)["alice"].net_code == 0
    warns = warnings_of(r, WarningKind.CROSS_CHECK)
    assert len(warns) == 1 and warns[0].scopes == ("kiểm tra chéo",)
    assert manifest_of(r.out_dir)["cross_check"]["ok"] is False
    assert [list(zip_files(z)) for z in r.end_zips] == [["app.py"]]


def test_merge_warnings(make_repo, cache_root: Path, tmp_path: Path) -> None:
    """TC-pipeline-20: file quá lớn không đổi → một cảnh báo, gộp phạm vi; tất định."""
    src = make_repo("repo")
    big = NL.join(f"big_value_{i:03d} = 'xxxxxxxxxxxxxxxxxxxxxxxx'" for i in range(40)) + NL
    src.commit("admin big", "2025-12-01T10:00:00+07:00", files={"big.py": big})
    _build_section14(src, False)
    repo = open_repo(src, cache_root)
    cfg = AppConfig(filters=FilterConfig(max_file_size_kb=1))

    a = analyze(repo, tmp_path / "a", config=cfg)
    big_a = [w for w in a.data.warnings if w.path == "big.py"]
    assert len(big_a) == 1 and big_a[0].kind == WarningKind.TOO_LARGE
    assert big_a[0].scopes == ("start", "end:all")

    b1 = analyze(repo, tmp_path / "b1", ("alice", "bob"), config=cfg, separate=True)
    big_b = [w for w in b1.data.warnings if w.path == "big.py"]
    assert len(big_b) == 1
    assert big_b[0].scopes == ("start", "end:multi-2users", "end:alice", "end:bob")
    b2 = analyze(repo, tmp_path / "b2", ("alice", "bob"), config=cfg, separate=True)
    assert b1.data.warnings == b2.data.warnings


def test_identity_warning(s14, tmp_path: Path) -> None:
    """TC-pipeline-21: GitHub API 403 → cảnh báo identity, path "", phạm vi định danh."""
    _, local = s14()
    repo = github_repo(local)
    resolver = make_resolver(repo.source, None, TOKEN, login_cache=tmp_path / "logins.json")
    assert isinstance(resolver.github, GitHubLoginResolver)
    resolver.github._session = FakeSession(FakeResponse(403))
    r = analyze(repo, tmp_path / "out", resolver=resolver, secrets=(TOKEN,))
    ident = warnings_of(r, WarningKind.IDENTITY)
    assert len(ident) == 1
    assert ident[0].path == "" and ident[0].scopes == ("định danh",)
    assert resolver.github.requests_made >= 1


def test_path_conflict(make_repo, cache_root: Path, tmp_path: Path) -> None:
    """TC-pipeline-22: A.py + a.py → path_conflict phạm vi start, end:all; vẫn đếm cả hai."""
    src = make_repo("repo")
    src.add_index_entry("100644", src.hash_blob(b"A = 1\n"), "A.py")
    src.add_index_entry("100644", src.hash_blob(b"a = 1\n"), "a.py")
    src.commit("admin case", "2025-12-01T10:00:00+07:00")
    _build_section14(src, False)
    r = analyze(open_repo(src, cache_root), tmp_path / "out")
    pc = warnings_of(r, WarningKind.PATH_CONFLICT)
    assert sorted(w.path for w in pc) == ["A.py", "a.py"]
    assert all(w.scopes == ("start", "end:all") for w in pc)
    for zip_path in (r.start_zip, *r.end_zips):
        assert {"A.py", "a.py"} <= zip_files(zip_path).keys()
    assert {"A.py", "a.py"} <= {f.path for f in r.data.files}


@pytest.mark.parametrize("users", [None, ("alice",)])
def test_windows_path(make_repo, cache_root: Path, tmp_path: Path,
                      users: tuple[str, ...] | None) -> None:
    """TC-pipeline-23: CON.py, a.py. → windows_path ở mọi phạm vi export; vẫn giữ và đếm."""
    src = make_repo("repo")
    src.git("config", "core.protectNTFS", "false")
    src.add_index_entry("100644", src.hash_blob(b"c = 1\n"), "CON.py")
    src.add_index_entry("100644", src.hash_blob(b"d = 1\n"), "a.py.")
    src.commit("admin win", "2025-12-01T10:00:00+07:00")
    _build_section14(src, False)
    r = analyze(open_repo(src, cache_root), tmp_path / "out", users,
                config=AppConfig(filters=FULL))
    end_scope = f"end:{r.data.exports[1].tag}"
    wp = {w.path: w for w in warnings_of(r, WarningKind.WINDOWS_PATH)}
    assert set(wp) == {"CON.py", "a.py."}
    assert all(w.scopes == ("start", end_scope) for w in wp.values())
    for zip_path in (r.start_zip, *r.end_zips):
        assert {"CON.py", "a.py."} <= zip_files(zip_path).keys()
    assert "CON.py" in {f.path for f in r.data.files}


# ---------- TC-24, TC-25: mandatory mục E ----------
@pytest.mark.parametrize("users", [None, ("alice",)])
def test_loc_counted_from_zip(s14, tmp_path: Path, users: tuple[str, ...] | None) -> None:
    """TC-pipeline-24: ExportSummary.loc = đếm lại từ zip giải nén (D36)."""
    _, repo = s14()
    r = analyze(repo, tmp_path / "out", users)
    for summary, zip_path in zip((r.data.exports), (r.start_zip, *r.end_zips), strict=True):
        dest = tmp_path / "unzipped" / zip_path.stem
        with zipfile.ZipFile(zip_path) as zf:
            zf.extractall(dest)
        dest.mkdir(parents=True, exist_ok=True)
        assert summary.loc == count_directory(dest, FilterConfig()).total_code
    assert r.data.exports[0].loc == r.data.loc_start
    assert r.data.exports[1].loc == r.data.loc_end


def _all_bytes(root: Path) -> list[tuple[str, bytes]]:
    """Nội dung mọi file dưới root, kèm nội dung đã giải nén của zip / xlsx."""
    out: list[tuple[str, bytes]] = []
    for p in sorted(root.rglob("*")):
        if not p.is_file():
            continue
        data = p.read_bytes()
        out.append((str(p), data))
        if p.suffix in (".zip", ".xlsx"):
            with zipfile.ZipFile(p) as zf:
                out += [(f"{p}!{n}", zf.read(n)) for n in zf.namelist()]
    return out


def test_token_not_leaked(s14, tmp_path: Path) -> None:
    """TC-pipeline-25: token / base64 không có trong mọi file đầu ra, run.log, cache."""
    _, local = s14()
    repo = github_repo(local)
    cache = tmp_path / "logins.json"
    resolver = make_resolver(repo.source, None, TOKEN, login_cache=cache)
    resolver.github._session = FakeSession(FakeResponse(403))
    out = tmp_path / "out"
    request = AnalysisRequest("main", jan(), AppConfig(), out)
    with run_log(out, secrets=(TOKEN,)):
        logging.getLogger("loc_counter.test").info("token=%s b64=%s", TOKEN, TOKEN_B64)
        with pytest.raises(GitError):
            repo.run(["rev-parse", "--verify", f"refs/heads/{TOKEN}"])
        r = run_analysis(repo, request, resolver, now=lambda: NOW)
    log = (out / "run.log").read_text(encoding="utf-8")
    assert "***" in log
    checked = _all_bytes(r.out_dir) + [("run.log", log.encode("utf-8")),
                                       ("config", (repo.git_dir / "config").read_bytes())]
    if cache.exists():
        checked.append(("logins", cache.read_bytes()))
    for name, data in checked:
        assert TOKEN.encode() not in data, name
        assert TOKEN_B64.encode() not in data, name
    for cls in (AnalysisRequest, PeriodContext, AnalysisResult):
        for f in dataclasses.fields(cls):
            assert not any(w in f.name.lower() for w in ("token", "secret", "auth", "password"))


# ---------- TC-26, TC-27: tiến độ, run.log ----------
def test_progress(s14, tmp_path: Path) -> None:
    """TC-pipeline-26: 5 bước đúng thứ tự; callback ném lỗi → ném lại, dọn thư mục tạm."""
    _, repo = s14()
    calls: list[tuple[Step, int, int]] = []
    analyze(repo, tmp_path / "out", progress=lambda s, i, n: calls.append((s, i, n)))
    order = [Step.EXPORT_START, Step.ATTRIBUTION, Step.EXPORT_END, Step.COUNT, Step.REPORT]
    assert calls == [(s, i, 5) for i, s in enumerate(order, 1)]
    assert [STEP_LABELS[s] for s in order] == [
        "Export start", "Gán tác giả", "Export end", "Đếm LOC", "Tạo báo cáo"]

    def failing(step: Step, i: int, n: int) -> None:
        if step == Step.EXPORT_END:
            raise RuntimeError("dừng")

    out2 = tmp_path / "out2"
    with pytest.raises(RuntimeError, match="dừng"):
        analyze(repo, out2, progress=failing)
    assert leftovers(out2) == [] and not (out2 / ALL_DIR).exists()


def test_run_log(s14, tmp_path: Path) -> None:
    """TC-pipeline-27: ghi đè, UTF-8, chỉ luồng hiện tại, gỡ handler và khôi phục level."""
    _, repo = s14()
    out = tmp_path / "out"
    pkg = logging.getLogger("loc_counter")
    root = logging.getLogger()
    old_root, old_pkg = root.level, pkg.level
    root.setLevel(logging.WARNING)
    pkg.setLevel(logging.NOTSET)
    try:
        with run_log(out) as path:
            logging.getLogger("loc_counter.test").info("lần một")
        assert path == (out / "run.log").resolve()

        def other() -> None:
            logging.getLogger("loc_counter.test").info("luồng khác ghi")

        with run_log(out):
            thread = threading.Thread(target=other)
            thread.start()
            thread.join()
            prepare_period(repo, "main", jan(), AccountResolver(IDS))
            logging.getLogger("loc_counter.test").info("tiếng Việt có dấu")
        text = (out / "run.log").read_bytes().decode("utf-8")
        assert "lần một" not in text and "luồng khác ghi" not in text
        assert "tiếng Việt có dấu" in text
        assert any(" INFO " in line and "C_START=" in line for line in text.splitlines())
        assert pkg.level == logging.NOTSET
        assert not [h for h in pkg.handlers if isinstance(h, logging.FileHandler)]
        (out / "run.log").unlink()
    finally:
        root.setLevel(old_root)
        pkg.setLevel(old_pkg)


def test_run_log_overlapping_threads(tmp_path: Path) -> None:
    """TC-pipeline-27: hai luồng lồng thời gian — luồng ra trước không làm mất log INFO (A18)."""
    pkg = logging.getLogger("loc_counter")
    root = logging.getLogger()
    old_root, old_pkg = root.level, pkg.level
    root.setLevel(logging.WARNING)
    pkg.setLevel(logging.NOTSET)
    a_in, b_in, a_out = threading.Event(), threading.Event(), threading.Event()
    errors: list[BaseException] = []

    def thread_a() -> None:
        try:
            with run_log(tmp_path / "a"):
                a_in.set()
                assert b_in.wait(10)
        except BaseException as exc:  # noqa: BLE001
            errors.append(exc)
        finally:
            a_out.set()

    def thread_b() -> None:
        try:
            assert a_in.wait(10)
            with run_log(tmp_path / "b"):
                b_in.set()
                assert a_out.wait(10)
                logging.getLogger("loc_counter.test").info("B ghi sau khi A ra")
        except BaseException as exc:  # noqa: BLE001
            errors.append(exc)

    try:
        threads = [threading.Thread(target=thread_a), threading.Thread(target=thread_b)]
        for t in threads:
            t.start()
        for t in threads:
            t.join(30)
        assert errors == []
        assert "B ghi sau khi A ra" in (tmp_path / "b" / "run.log").read_text(encoding="utf-8")
        assert pkg.level == logging.NOTSET
    finally:
        root.setLevel(old_root)
        pkg.setLevel(old_pkg)


# ---------- TC-28: tên thư mục ----------
def test_result_dir_name() -> None:
    """TC-pipeline-28: <repo>_<branch>_<start>_<end>_<tag>, cắt phần dài, giờ từng phía."""
    p = jan()
    for source, repo in (("https://github.com/org/repo.git", "repo"),
                         ("git@github.com:org/Repo.GIT", "Repo"),
                         ("file:///D:/x/r", "r"), ("D:\\x\\r\\", "r"), ("/home/u/r/.git", "r")):
        assert result_dir_name(source, "main", p, "all") == (
            f"{repo}_main_2026-01-01_2026-01-31_all"), source
    assert result_dir_name("r", "feature/x", p, "alice") == (
        "r_feature_x_2026-01-01_2026-01-31_alice")

    long = "a" * 50
    cut = "a" * 33 + "-" + hashlib.sha1(long.encode()).hexdigest()[:6]
    assert len(cut) == 40
    assert result_dir_name(f"https://h/o/{long}", long, p, "all") == (
        f"{cut}_{cut}_2026-01-01_2026-01-31_all")

    timed = make_period(*JAN, start_time=time(8, 30))
    assert result_dir_name("r", "main", timed, "all") == "r_main_2026-01-01T0830_2026-01-31_all"
    end_2359 = make_period(*JAN, end_time=time(23, 59))
    assert result_dir_name("r", "main", end_2359, "all") == "r_main_2026-01-01_2026-01-31_all"
    end_1730 = make_period(*JAN, end_time=time(17, 30))
    assert result_dir_name("r", "main", end_1730, "all") == (
        "r_main_2026-01-01_2026-01-31T1730_all")


# ---------- TC-29 .. TC-35 ----------
class NullLoginSession:
    """Session giả luôn trả 200 không có login; đếm request."""

    def __init__(self) -> None:
        self.calls = 0

    def get(self, url: str, headers: dict, timeout: float, auth=None) -> FakeResponse:  # noqa: ANN001
        self.calls += 1
        return FakeResponse(200, {"author": None})


def test_shared_resolver(s14, tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    """TC-pipeline-29: cùng một resolver cho danh sách user và gán tác giả; mỗi email 1 request."""
    _, local = s14()
    repo = github_repo(local)
    resolver = make_resolver(repo.source, None, TOKEN, login_cache=tmp_path / "logins.json")
    session = NullLoginSession()
    resolver.github._session = session
    users_spy = Spy(monkeypatch, pipeline, "summarize_users")
    attr_spy = Spy(monkeypatch, pipeline, "attribute_period")
    analyze(repo, tmp_path / "out", resolver=resolver)
    assert users_spy.calls[0][0][1] is resolver
    assert attr_spy.calls[0][0][3] is resolver
    assert session.calls == 2  # alice@example.com, bob@example.com


def test_prepare_period_context(s14, tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    """TC-pipeline-30: run_analysis(context=…) không tìm lại snapshot / commit; khác nhánh → lỗi."""
    _, repo = s14()
    resolver = AccountResolver(IDS)
    ctx = prepare_period(repo, "main", jan(), resolver)
    assert [u.account for u in ctx.users.users] == ["alice", "bob"]
    snap_spy = Spy(monkeypatch, pipeline, "find_snapshots")
    log_spy = Spy(monkeypatch, repo, "log_commits")
    r = analyze(repo, tmp_path / "out", resolver=resolver, context=ctx)
    assert snap_spy.calls == [] and log_spy.calls == []
    assert (r.data.loc_start, r.data.loc_end) == (10, 20)

    with pytest.raises(ValueError):
        analyze(repo, tmp_path / "o2", resolver=resolver,
                context=dataclasses.replace(ctx, branch="dev"))
    with pytest.raises(ValueError):
        analyze(repo, tmp_path / "o3", resolver=resolver,
                context=dataclasses.replace(ctx, period=make_period(date(2026, 1, 2), JAN[1])))


def test_analysis_result(s14, tmp_path: Path) -> None:
    """TC-pipeline-31: đường dẫn trong AnalysisResult tồn tại, khớp data và manifest."""
    _, repo = s14()
    r = analyze(repo, tmp_path / "out", ("alice", "bob"), separate=True)
    for p in (r.start_zip, *r.end_zips, *r.report_files):
        assert p.is_file() and p.parent == r.out_dir
    assert [p.name for p in r.end_zips] == [e.file_name for e in r.data.exports[1:]]
    assert r.start_zip.name == r.data.exports[0].file_name
    assert [p.name for p in r.report_files] == REPORT_NAMES
    assert sorted(u.account for u in r.period_users) == ["alice", "bob"]
    assert r.generated_at == NOW
    m = manifest_of(r.out_dir)
    assert m["totals"] == {"loc_start": 10, "loc_end": 20, "loc_delta": 10,
                           "added_code": 13, "deleted_code": 3, "net_code": 10,
                           "reported_loc_end": 20, "reported_loc_delta": 10}
    assert [(e["file"], e["loc"]) for e in m["exports"]] == [
        (e.file_name, e.loc) for e in r.data.exports]


def test_timezone_from_period(s14, tmp_path: Path) -> None:
    """TC-pipeline-32: múi giờ của period thắng config; manifest và timestamp zip theo UTC."""
    _, repo = s14()
    period = make_period(*JAN, timezone="UTC")
    r = analyze(repo, tmp_path / "out", period=period,
                config=AppConfig(timezone="Asia/Bangkok"))
    m = manifest_of(r.out_dir)
    assert m["config"]["timezone"] == "UTC" == m["timezone"]
    with zipfile.ZipFile(r.start_zip) as zf:
        assert {i.date_time for i in zf.infolist()} == {(2025, 12, 15, 3, 0, 0)}
    with zipfile.ZipFile(r.end_zips[0]) as zf:
        assert {i.date_time for i in zf.infolist()} == {(2026, 1, 8, 3, 0, 0)}


def test_path_length_check(s14, tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    """TC-pipeline-33: Windows chưa bật long path + out_root dài → lỗi trước việc nặng."""
    _, repo = s14()
    out = tmp_path / ("o" * max(1, 230 - len(str(tmp_path)) - 1))
    monkeypatch.setattr(pipeline, "_path_limit_applies", lambda: True)
    forbid(monkeypatch, "attribute_period", "plan_export")
    request = AnalysisRequest("main", jan(), AppConfig(), out)
    with pytest.raises(PipelineError, match="Đường dẫn kết quả quá dài") as exc:
        run_analysis(repo, request, AccountResolver(IDS))
    assert "259" in str(exc.value)
    assert leftovers(out) == []

    monkeypatch.setattr(pipeline, "_path_limit_applies", lambda: False)
    pipeline._check_path_lengths(out, out / ("x" * 100), ["y" * 100])


class FakeWinreg:
    """winreg giả: QueryValueEx trả `value`, hoặc OpenKey ném OSError khi value là None."""

    HKEY_LOCAL_MACHINE = object()

    def __init__(self, value: int | None) -> None:
        self.value = value

    def OpenKey(self, root: object, path: str) -> contextlib.AbstractContextManager[object]:  # noqa: N802
        if self.value is None:
            raise OSError("không đọc được registry")
        assert root is self.HKEY_LOCAL_MACHINE
        assert path == r"SYSTEM\CurrentControlSet\Control\FileSystem"
        return contextlib.nullcontext(object())

    def QueryValueEx(self, key: object, name: str) -> tuple[int, int]:  # noqa: N802
        assert name == "LongPathsEnabled"
        return self.value, 4  # type: ignore[return-value]


def test_long_paths_registry(monkeypatch: pytest.MonkeyPatch) -> None:
    """TC-pipeline-33: đọc LongPathsEnabled (1 → bật, 0 / lỗi → chưa bật); không phải nt."""
    for value, enabled in ((1, True), (0, False), (None, False)):
        monkeypatch.setitem(sys.modules, "winreg", FakeWinreg(value))
        assert pipeline._long_paths_enabled() is enabled
        with monkeypatch.context() as m:
            m.setattr(pipeline.os, "name", "nt")
            assert pipeline._path_limit_applies() is not enabled
    with monkeypatch.context() as m:
        m.setattr(pipeline.os, "name", "posix")
        assert pipeline._path_limit_applies() is False


# ---------- TC-36: nhánh lỗi hệ thống tệp ----------
def _fail_when(orig: Callable[..., Any], prefix: str, exc: BaseException) -> Callable[..., Any]:
    """Bọc hàm (path đầu tiên là tham số đầu): ném exc khi tên path bắt đầu bằng prefix."""

    def wrapper(path: Any, *args: Any, **kwargs: Any) -> Any:
        if Path(path).name.startswith(prefix):
            raise exc
        return orig(path, *args, **kwargs)

    return wrapper


def test_filesystem_error_branches(s14, tmp_path: Path, monkeypatch: pytest.MonkeyPatch,
                                   caplog: pytest.LogCaptureFixture) -> None:
    """TC-pipeline-36: lỗi mkdir / FileHandler / os.replace / rmtree → đúng thông báo mục 7."""
    _, repo = s14()
    with monkeypatch.context() as m:
        m.setattr(Path, "mkdir", _fail_when(Path.mkdir, "new", OSError("giả")))
        with pytest.raises(PipelineError, match="Không tạo được thư mục đầu ra"):
            pipeline._prepare_out_root(tmp_path / "new")
    with monkeypatch.context() as m:
        m.setattr(logging, "FileHandler", _fail_when(logging.FileHandler, "run", OSError("x")))
        with (pytest.raises(PipelineError, match="Không ghi được run.log"),
              run_log(tmp_path / "log")):
            pass

    out = tmp_path / "out"
    out.mkdir()
    with monkeypatch.context() as m:
        m.setattr(Path, "mkdir", _fail_when(Path.mkdir, pipeline.TMP_PREFIX, OSError("giả")))
        with pytest.raises(PipelineError, match="Không tạo được thư mục tạm trong"):
            analyze(repo, out)
    assert sorted(p.name for p in out.iterdir()) == ["run.log"]

    with monkeypatch.context() as m:
        m.setattr(os, "replace", _fail_when(os.replace, pipeline.TMP_PREFIX, OSError("giả")))
        with pytest.raises(PipelineError, match="Không đổi tên thư mục tạm thành"):
            analyze(repo, out)
    assert leftovers(out) == [] and not (out / ALL_DIR).exists()

    first = analyze(repo, out)
    before = digests(first.out_dir)
    cases = ((OSError("giả"), PipelineError), (KeyboardInterrupt(), KeyboardInterrupt))
    for exc, expected in cases:
        with monkeypatch.context() as m:
            m.setattr(os, "replace", _fail_when(os.replace, pipeline.TMP_PREFIX, exc))
            with pytest.raises(expected) as info:
                analyze(repo, out, now=NOW + timedelta(days=1))
        if expected is PipelineError:
            assert "Không thay thế được thư mục kết quả cũ" in str(info.value)
        assert digests(first.out_dir) == before
        assert leftovers(out) == []

    caplog.set_level(logging.WARNING, logger="loc_counter")
    replace = os.replace

    def replace_tmp_and_old_fail(src: Any, dst: Any) -> None:
        if Path(src).name.startswith((pipeline.TMP_PREFIX, pipeline.OLD_PREFIX)):
            raise OSError("giả")
        replace(src, dst)

    with monkeypatch.context() as m:
        m.setattr(os, "replace", replace_tmp_and_old_fail)
        with pytest.raises(PipelineError) as info:
            analyze(repo, out, now=NOW + timedelta(days=1))
    lost = [n for n in leftovers(out) if n.startswith(pipeline.OLD_PREFIX)]
    assert len(lost) == 1 and not (out / ALL_DIR).exists()
    assert f"dữ liệu cũ nằm ở {out.resolve() / lost[0]}" in str(info.value)
    assert any(r.levelno == logging.ERROR and lost[0] in r.getMessage() for r in caplog.records)
    assert digests(out / lost[0]) == before
    os.replace(out / lost[0], out / ALL_DIR)
    caplog.clear()

    with monkeypatch.context() as m:
        m.setattr(shutil, "rmtree", _fail_when(shutil.rmtree, pipeline.OLD_PREFIX, OSError("x")))
        again = analyze(repo, out, now=NOW + timedelta(days=2))
    assert manifest_of(again.out_dir)["generated_at"] == (NOW + timedelta(days=2)).isoformat()
    assert any("Không xóa được thư mục cũ" in r.getMessage() for r in caplog.records)
    old = [n for n in leftovers(out) if n.startswith(pipeline.OLD_PREFIX)]
    assert len(old) == 1
    shutil.rmtree(out / old[0])


def test_zip_timestamps(s14, tmp_path: Path) -> None:
    """TC-pipeline-34: entry zip start mang committer date C_START, end mang C_END (Bangkok)."""
    _, repo = s14()
    r = analyze(repo, tmp_path / "out")
    with zipfile.ZipFile(r.start_zip) as zf:
        assert {i.date_time for i in zf.infolist()} == {(2025, 12, 15, 10, 0, 0)}
    with zipfile.ZipFile(r.end_zips[0]) as zf:
        assert {i.date_time for i in zf.infolist()} == {(2026, 1, 8, 10, 0, 0)}


def test_make_resolver(tmp_path: Path) -> None:
    """TC-pipeline-35: GitHub + token → GitHubLoginResolver; không token / local → None."""
    cache = tmp_path / "logins.json"
    r = make_resolver("https://github.com/org/repo", IDS, TOKEN, login_cache=cache)
    assert isinstance(r.github, GitHubLoginResolver) and r.github._cache_path == cache
    assert r.identities is IDS
    assert make_resolver("https://github.com/org/repo", None, None).github is None
    assert make_resolver(str(tmp_path), None, TOKEN).github is None


# ---------- TC-38, 39: dùng lại kết quả gán tác giả, log luồng con (A21, A22) ----------
def test_attribution_memo(s14, tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    """TC-pipeline-38: cùng ctx + filters + resolver → không gán tác giả lại; khác → tính lại."""
    _, repo = s14()
    resolver = AccountResolver(IDS)
    ctx = prepare_period(repo, "main", jan(), resolver)
    calls: list[FilterConfig] = []
    real = pipeline.attribute_period

    def spy(*args: Any, **kwargs: Any) -> Any:
        calls.append(args[4])
        return real(*args, **kwargs)

    monkeypatch.setattr(pipeline, "attribute_period", spy)
    analyze(repo, tmp_path / "o1", resolver=resolver, context=ctx)
    assert len(calls) == 1
    r2 = analyze(repo, tmp_path / "o2", ("alice",), resolver=resolver, context=ctx)
    assert len(calls) == 1
    log2 = (tmp_path / "o2" / "run.log").read_text(encoding="utf-8")
    assert "Dùng lại kết quả gán tác giả đã tính cho kỳ này" in log2

    fresh_resolver = AccountResolver(IDS)
    fresh = prepare_period(repo, "main", jan(), fresh_resolver)
    r2b = analyze(repo, tmp_path / "o2b", ("alice",), resolver=fresh_resolver, context=fresh)
    assert len(calls) == 2
    assert digests(r2.out_dir) == digests(r2b.out_dir)  # cùng now → trùng cả manifest
    assert r2.data == r2b.data

    docstrings = AppConfig(FilterConfig(count_docstrings=True))
    analyze(repo, tmp_path / "o3", resolver=resolver, context=ctx, config=docstrings)
    assert len(calls) == 3 and calls[-1].count_docstrings
    analyze(repo, tmp_path / "o4", resolver=AccountResolver(IDS), context=ctx, config=docstrings)
    assert len(calls) == 4
    assert len(ctx.attribution_memo) == 1
    for name in ("o1", "o3", "o4"):
        text = (tmp_path / name / "run.log").read_text(encoding="utf-8")
        assert "Dùng lại kết quả gán tác giả" not in text

    plain = PeriodContext(ctx.branch, ctx.period, ctx.snapshots, ctx.commits, ctx.users)
    assert plain == ctx
    assert plain.attribution_memo == {}
    assert dataclasses.replace(ctx).attribution_memo == {}  # replace không chia sẻ memo

    # memo chỉ giữ weakref: resolver bị thu hồi thì không còn được giữ (không giữ token)
    import gc
    import weakref

    temp = AccountResolver(IDS)
    analyze(repo, tmp_path / "o5", resolver=temp, context=ctx, config=docstrings)
    assert len(calls) == 5
    probe = weakref.ref(temp)
    del temp
    gc.collect()
    assert probe() is None


def _warn_in_pool(message: str) -> None:
    import contextvars
    from concurrent.futures import ThreadPoolExecutor

    log = logging.getLogger("loc_counter.test")
    with ThreadPoolExecutor(2) as pool:
        pool.submit(contextvars.copy_context().run, log.warning, message).result()


def test_run_log_worker_threads(tmp_path: Path) -> None:
    """TC-pipeline-39: (a) luồng con có context vào run.log; (b) Thread thường thì không; (d)."""
    with run_log(tmp_path / "x"):
        _warn_in_pool("con có context")
        other = threading.Thread(target=logging.getLogger("loc_counter.test").warning,
                                 args=("luồng thường",))
        other.start()
        other.join()
    text = (tmp_path / "x" / "run.log").read_text(encoding="utf-8")
    assert "con có context" in text
    assert "luồng thường" not in text
    assert pipeline._RUN_LOG_ID.get(None) is None

    a_in, b_in, a_done = threading.Event(), threading.Event(), threading.Event()
    errors: list[BaseException] = []
    after: list[str | None] = []

    def session(name: str, mine: threading.Event, theirs: threading.Event) -> None:
        try:
            with run_log(tmp_path / name):
                mine.set()
                assert theirs.wait(10)
                _warn_in_pool(f"từ phiên {name}")
                if name == "a":
                    a_done.set()
                else:
                    assert a_done.wait(10)
            after.append(pipeline._RUN_LOG_ID.get(None))
        except BaseException as exc:  # noqa: BLE001
            errors.append(exc)

    threads = [threading.Thread(target=session, args=("a", a_in, b_in)),
               threading.Thread(target=session, args=("b", b_in, a_in))]
    for t in threads:
        t.start()
    for t in threads:
        t.join(30)
    assert errors == []
    assert after == [None, None]
    log_a = (tmp_path / "a" / "run.log").read_text(encoding="utf-8")
    log_b = (tmp_path / "b" / "run.log").read_text(encoding="utf-8")
    assert "từ phiên a" in log_a and "từ phiên b" not in log_a
    assert "từ phiên b" in log_b and "từ phiên a" not in log_b


def test_run_log_exit_other_context(tmp_path: Path) -> None:
    """TC-pipeline-39: (e) vào with ở một context, thoát ở context khác → vẫn dọn sạch."""
    import contextvars

    pkg = logging.getLogger("loc_counter")
    old_level = pkg.level
    cm = run_log(tmp_path / "e")
    contextvars.copy_context().run(cm.__enter__)
    cm.__exit__(None, None, None)  # reset token ở context khác: không được làm sót dọn dẹp
    assert not [h for h in pkg.handlers if isinstance(h, logging.FileHandler)]
    assert pkg.level == old_level
    assert pipeline._level_sessions == 0
    (tmp_path / "e" / "run.log").unlink()  # Windows: handler đã đóng nên xóa được


def test_run_log_attribution_workers(make_repo, cache_root: Path, tmp_path: Path,
                                     monkeypatch: pytest.MonkeyPatch) -> None:
    """TC-pipeline-39: (c) WARNING phát trong luồng gán tác giả song song có trong run.log."""
    from test_attribution import build_merge_parallel, prepare

    from loc_counter import attribution

    src = make_repo()
    build_merge_parallel(src)
    src.commit("alice extra", "2026-01-08T10:00:00+07:00", author=ALICE,
               files={"extra.py": "x = 1\n"})
    ctx = prepare(src, cache_root)
    monkeypatch.setattr(attribution, "MAX_MERGE_DEPTH", 0)
    monkeypatch.setattr(attribution, "ATTRIBUTION_WORKERS", 4)
    seen: list[str] = []
    real = ctx.repo.blame

    def spy(*args: Any, **kwargs: Any) -> Any:
        seen.append(threading.current_thread().name)
        return real(*args, **kwargs)

    ctx.repo.blame = spy  # type: ignore[method-assign]
    with run_log(tmp_path / "out"):
        ctx.attribute()
    text = (tmp_path / "out" / "run.log").read_text(encoding="utf-8")
    assert "Vượt giới hạn đệ quy merge" in text
    assert any(name.startswith("loc-attr") for name in seen)


# ---------- TC-40: chế độ gán tác giả last_commit ----------
def test_attribution_mode_last_commit(s14, tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    """TC-pipeline-40: last_commit — tất cả 10/20; alice 10/10; bob 10/20; memo theo chế độ."""
    _, repo = s14()

    def run(out: str, users: tuple[str, ...] | None, mode: str,
            ctx: PeriodContext | None = None) -> AnalysisResult:
        request = AnalysisRequest("main", jan(), AppConfig(), tmp_path / out, users,
                                  attribution_mode=mode)  # type: ignore[arg-type]
        resolver = AccountResolver(IDS)
        with run_log(tmp_path / out):
            return run_analysis(repo, request, resolver, context=ctx, now=lambda: NOW)

    r = run("all", None, "last_commit")
    assert (r.data.loc_start, r.data.loc_end) == (10, 20)
    assert r.data.cross_check.diff_percent == 0 and not warnings_of(r, WarningKind.CROSS_CHECK)
    assert r.data.options.attribution_mode == "last_commit"
    assert manifest_of(r.out_dir)["options"]["attribution_mode"] == "last_commit"
    rows = rows_of(r)
    assert (rows["alice"].added_code, rows["alice"].deleted_code) == (0, 0)
    assert (rows["bob"].added_code, rows["bob"].deleted_code) == (13, 3)
    assert (run("a", ("alice",), "last_commit").data.loc_end,
            run("b", ("bob",), "last_commit").data.loc_end) == (10, 20)
    assert manifest_of(run("c", None, "line").out_dir)["options"]["attribution_mode"] == "line"

    ctx = prepare_period(repo, "main", jan(), AccountResolver(IDS))
    spy = Spy(monkeypatch, pipeline, "attribute_period")
    resolver = AccountResolver(IDS)
    request = AnalysisRequest("main", jan(), AppConfig(), tmp_path / "m")
    with run_log(tmp_path / "m"):
        run_analysis(repo, request, resolver, context=ctx, now=lambda: NOW)
        run_analysis(repo, request, resolver, context=ctx, now=lambda: NOW)
        assert len(spy.calls) == 1
        last = dataclasses.replace(request, attribution_mode="last_commit")
        run_analysis(repo, last, resolver, context=ctx, now=lambda: NOW)
    assert len(spy.calls) == 2 and spy.calls[1][1]["mode"] == "last_commit"
    with pytest.raises(ValueError, match="attribution_mode"):
        run("bad", None, "x")


# ---------- TC-41: cách tính LOC end ----------
def test_loc_end_mode(s14, tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    """TC-pipeline-41: loc_end_mode — net_code mặc định; added_code = 10 + 13; export không đổi."""
    _, repo = s14()
    ctx = prepare_period(repo, "main", jan(), AccountResolver(IDS))
    resolver = AccountResolver(IDS)
    spy = Spy(monkeypatch, pipeline, "attribute_period")

    def run(out: str, mode: str | None) -> AnalysisResult:
        extra = {} if mode is None else {"loc_end_mode": mode}
        request = AnalysisRequest("main", jan(), AppConfig(), tmp_path / out, **extra)
        with run_log(tmp_path / out):
            return run_analysis(repo, request, resolver, context=ctx, now=lambda: NOW)

    net = run("net", None)
    assert net.data.options.loc_end_mode == "net_code"
    assert (net.data.loc_end, net.data.reported_loc_end) == (20, 20)
    added = run("added", "added_code")
    assert len(spy.calls) == 1  # chỉ đổi cách tính LOC end → dùng lại kết quả gán tác giả
    assert added.data.options.loc_end_mode == "added_code"
    assert (added.data.loc_end, added.data.reported_loc_end) == (20, 23)
    assert added.data.cross_check == net.data.cross_check and added.data.cross_check.ok
    m = manifest_of(added.out_dir)
    assert (m["totals"]["loc_end"], m["totals"]["reported_loc_end"],
            m["totals"]["reported_loc_delta"]) == (20, 23, 13)
    assert m["options"]["loc_end_mode"] == "added_code"
    for name in [*(p.name for p in net.end_zips), net.start_zip.name, "loc_by_user.csv",
                 "loc_by_language.csv", "loc_by_file.csv", "loc_by_pr.csv", "warnings.csv"]:
        assert (hashlib.sha256((net.out_dir / name).read_bytes()).digest()
                == hashlib.sha256((added.out_dir / name).read_bytes()).digest()), name
    log = (tmp_path / "added" / "run.log").read_text(encoding="utf-8")
    assert "LOC end theo code thêm: 10 + 13 = 23 (đếm trên export: 20)" in log
    assert "LOC end theo code thêm" not in (tmp_path / "net" / "run.log").read_text(
        encoding="utf-8")
    with pytest.raises(ValueError, match="loc_end_mode"):
        run("bad", "x")
    assert not (tmp_path / "bad").exists() or not any(
        p.is_dir() for p in (tmp_path / "bad").iterdir())


def test_pr_stats_end_to_end(make_repo, cache_root: Path, tmp_path: Path) -> None:
    """TC-pipeline-42: LOC theo PR — merge PR, squash, trực tiếp; file Theo_PR / loc_by_pr.csv."""
    bob = ("bob", "bob@example.com")
    src = make_repo()
    base = src.commit("base", "2025-12-20T10:00:00+07:00", files={"app.py": "x0 = 0\n"})
    src.git("checkout", "-q", "-b", "feat", base)
    src.commit("bob b", "2026-01-03T10:00:00+07:00", author=bob,
               files={"b.py": NL.join(f"b{i} = {i}" for i in range(4)) + NL})
    src.commit("alice c", "2026-01-04T10:00:00+07:00", author=ALICE,
               files={"c.py": "c1 = 1" + NL + "c2 = 2" + NL})
    src.checkout("main")
    src.commit("alice main", "2026-01-05T10:00:00+07:00", author=ALICE,
               files={"app.py": NL.join(["x0 = 0", "a1 = 1", "a2 = 2", "a3 = 3"]) + NL})
    src.merge("feat", "Merge pull request #7 from o/feat\n\nThêm b.py",
              "2026-01-06T10:00:00+07:00", author=("carol", "carol@example.com"))
    src.commit("Sửa lỗi (#9)", "2026-01-07T10:00:00+07:00", author=ALICE,
               files={"d.py": "d1 = 1" + NL})

    r = analyze(open_repo(src, cache_root), tmp_path / "out")
    assert [(p.pr, p.account, p.commits, p.added_code, p.deleted_code)
            for p in r.data.pr_rows] == [
        (7, "alice", 1, 2, 0), (7, "bob", 1, 4, 0), (9, "alice", 1, 1, 0), (None, "alice", 1, 3, 0)]
    assert "loc_by_pr.csv" in [p.name for p in r.report_files]
    ws = load_workbook(r.out_dir / "loc_report.xlsx")["Theo_PR"]
    sheet = list(ws.iter_rows(values_only=True))
    assert len(sheet) == 5 and sheet[1][:5] == (7, "Thêm b.py", sheet[1][2], "2026-01-06", "alice")
    for u in r.data.user_rows:      # bất biến: Σ theo PR = số liệu theo user
        mine = [p for p in r.data.pr_rows if p.account == u.account]
        assert (sum(p.commits for p in mine), sum(p.added_code for p in mine),
                sum(p.deleted_code for p in mine)) == (u.commits, u.added_code, u.deleted_code)
    again = analyze(open_repo(src, cache_root), tmp_path / "out2")
    assert digests(again.out_dir)["loc_by_pr.csv"] == digests(r.out_dir)["loc_by_pr.csv"]
