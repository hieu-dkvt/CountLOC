"""Báo cáo của một lần phân tích: xlsx, CSV, warnings.csv, manifest (§12, D40)."""

from __future__ import annotations

import csv
import io
import json
import logging
import re
import zipfile
from collections.abc import Callable, Iterable, Mapping, Sequence
from dataclasses import dataclass
from datetime import UTC, datetime
from pathlib import Path
from typing import TYPE_CHECKING, Any, Literal
from zoneinfo import ZoneInfo

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from loc_counter import __version__
from loc_counter.config import AppConfig, LocCounterError, config_to_dict
from loc_counter.counter import (
    CrossCheck,
    FileRow,
    LanguageRow,
    UserMetrics,
    WarningKind,
    WeekRow,
    total_metrics,
)

if TYPE_CHECKING:
    from openpyxl.worksheet.worksheet import Worksheet

    from loc_counter.git_repo import CommitInfo
    from loc_counter.snapshot import Period

logger = logging.getLogger(__name__)

XLSX_NAME = "loc_report.xlsx"
USER_CSV_NAME = "loc_by_user.csv"
LANGUAGE_CSV_NAME = "loc_by_language.csv"
FILE_CSV_NAME = "loc_by_file.csv"
WARNINGS_CSV_NAME = "warnings.csv"
MANIFEST_NAME = "manifest.json"
SHEET_NAMES = ("Tong_quan", "Theo_ngon_ngu", "Theo_user", "Theo_file", "Theo_tuan", "Canh_bao")
XLSX_MAX_CELL_CHARS = 32767

SCOPE_START = "start"
SCOPE_ATTRIBUTION = "gán tác giả"
SCOPE_CROSS_CHECK = "kiểm tra chéo"
SCOPE_IDENTITY = "định danh"

_NOTE_USER_EXPORT = ("Bản export end theo user dùng để kiểm tra và đếm LOC, "
                    "có thể không build/chạy được.")

_SURROGATE_RE = re.compile("[\ud800-\udfff]")
_ILLEGAL_RE = re.compile("[\x00-\x08\x0b\x0c\x0e-\x1f\ufffe\uffff]")
_CSV_DANGEROUS = ("=", "+", "-", "@", "\t", "\r")
_CORE_TIME_RE = re.compile(r"(<dcterms:(created|modified)\b[^>]*>)[^<]*(</dcterms:\2>)")
_ZIP_EPOCH = datetime(1980, 1, 1)

_INT = "#,##0"
_DELTA = "+#,##0;-#,##0;0"
_PERCENT = "0.000"

ExportRole = Literal["start", "end", "separate"]


def end_scope(tag: str) -> str:
    """Phạm vi của export end: "end:<tag>"."""
    return f"end:{tag}"


class ReportError(LocCounterError):
    """Lỗi ghi báo cáo (thông báo tiếng Việt)."""


@dataclass(frozen=True)
class ReportWarning:
    """Một cảnh báo đã gộp phạm vi (một hàng của warnings.csv)."""

    kind: WarningKind
    path: str
    message: str
    scopes: tuple[str, ...]


@dataclass(frozen=True)
class ExportSummary:
    """Một zip đã ghi: start, end (gộp / all) hoặc zip riêng từng user."""

    role: ExportRole
    file_name: str
    tag: str | None
    accounts: tuple[str, ...] | None
    files: int
    loc: int
    cross_check: CrossCheck | None = None


@dataclass(frozen=True)
class ReportOptions:
    """Tùy chọn người dùng chọn (ghi vào manifest)."""

    selected_users: tuple[str, ...] | None = None
    exclude_bots: bool = False
    separate_users: bool = False


@dataclass(frozen=True)
class ReportData:
    """Toàn bộ dữ liệu đã tính sẵn của một lần phân tích (pipeline dựng)."""

    repo: str
    branch: str
    period: Period
    c_start: CommitInfo | None
    c_end: CommitInfo
    users: tuple[str, ...]
    all_users: bool
    config: AppConfig
    options: ReportOptions
    loc_start: int
    loc_end: int
    languages: tuple[LanguageRow, ...]
    user_rows: tuple[UserMetrics, ...]
    files: tuple[FileRow, ...]
    weeks: tuple[WeekRow, ...]
    cross_check: CrossCheck
    warnings: tuple[ReportWarning, ...] = ()
    exports: tuple[ExportSummary, ...] = ()
    tool_version: str = __version__

    def __post_init__(self) -> None:
        if (self.cross_check.loc_start != self.loc_start
                or self.cross_check.loc_end != self.loc_end):
            raise ValueError("cross_check không khớp loc_start / loc_end của ReportData")

    @property
    def loc_delta(self) -> int:
        """loc_end - loc_start."""
        return self.loc_end - self.loc_start


# ---------- chuẩn hóa ----------
def clean_text(value: str) -> str:
    """Thay surrogate lẻ bằng U+FFFD, bỏ ký tự điều khiển XML không cho phép (giữ \\t \\n \\r)."""
    return _ILLEGAL_RE.sub("", _SURROGATE_RE.sub("\ufffd", value))


def _sorted_languages(rows: Iterable[LanguageRow]) -> list[LanguageRow]:
    return sorted(rows, key=lambda r: (-r.loc_end, -r.loc_start, r.language))


def _sorted_users(rows: Iterable[UserMetrics]) -> list[UserMetrics]:
    return sorted(rows, key=lambda r: (-r.commits, r.account))


def _sorted_files(rows: Iterable[FileRow]) -> list[FileRow]:
    return sorted(rows, key=lambda r: r.path)


def _sorted_weeks(rows: Iterable[WeekRow]) -> list[WeekRow]:
    return sorted(rows, key=lambda r: (r.week_start, r.account))


def _sorted_warnings(rows: Iterable[ReportWarning]) -> list[ReportWarning]:
    return sorted(rows, key=lambda w: (w.path, w.kind.value))


# Cột: (tiêu đề, hàm lấy giá trị, định dạng số)
_Column = tuple[str, Callable[[Any], Any], str | None]

_LANGUAGE_COLUMNS: tuple[_Column, ...] = (
    ("Ngôn ngữ", lambda r: r.language, None),
    ("Số file start", lambda r: r.files_start, _INT),
    ("LOC start", lambda r: r.loc_start, _INT),
    ("Số file end", lambda r: r.files_end, _INT),
    ("LOC end", lambda r: r.loc_end, _INT),
    ("Chênh lệch", lambda r: r.delta, _DELTA),
)
_USER_COLUMNS: tuple[_Column, ...] = (
    ("User", lambda r: r.account, None),
    ("Bot", lambda r: "có" if r.is_bot else "không", None),
    ("Commit", lambda r: r.commits, _INT),
    ("Code thêm", lambda r: r.added_code, _INT),
    ("Code xóa", lambda r: r.deleted_code, _INT),
    ("Code ròng", lambda r: r.net_code, _DELTA),
    ("File thay đổi", lambda r: r.files_changed, _INT),
    ("Dòng thêm", lambda r: r.added_lines, _INT),
    ("Dòng xóa", lambda r: r.deleted_lines, _INT),
)
_FILE_COLUMNS: tuple[_Column, ...] = (
    ("Đường dẫn", lambda r: r.path, None),
    ("Đổi tên từ", lambda r: r.renamed_from or "", None),
    ("Ngôn ngữ", lambda r: r.language, None),
    ("LOC start", lambda r: r.loc_start, _INT),
    ("LOC end", lambda r: r.loc_end, _INT),
    ("Chênh lệch", lambda r: r.delta, _DELTA),
)
_WEEK_COLUMNS: tuple[_Column, ...] = (
    ("Tuần", lambda r: r.week, None),
    ("Ngày bắt đầu", lambda r: r.week_start.isoformat(), None),
    ("User", lambda r: r.account, None),
    ("Commit", lambda r: r.commits, _INT),
    ("Code thêm", lambda r: r.added_code, _INT),
    ("Code xóa", lambda r: r.deleted_code, _INT),
    ("Code ròng", lambda r: r.net_code, _DELTA),
)
_WARNING_COLUMNS: tuple[_Column, ...] = (
    ("Loại", lambda w: w.kind.value, None),
    ("Đường dẫn", lambda w: w.path, None),
    ("Nội dung", lambda w: w.message, None),
    ("Phạm vi", lambda w: "; ".join(w.scopes), None),
)


# ---------- ghi file ----------
def _write_bytes(path: Path, data: bytes) -> Path:
    """Ghi bytes; OSError → ReportError tiếng Việt."""
    try:
        path.write_bytes(data)
    except OSError as exc:
        logger.error("Không ghi được file báo cáo %s: %s", path, exc)
        raise ReportError(f"Không ghi được file báo cáo {path}: {exc}") from exc
    return path


# ---------- xlsx ----------
def _put(ws: Worksheet, row: int, col: int, value: Any, fmt: str | None = None,
         bold: bool = False) -> None:
    """Ghi một ô: văn bản luôn kiểu chuỗi (đã làm sạch, cắt độ dài); chuỗi rỗng → ô trống."""
    if isinstance(value, str):
        value = clean_text(value).replace("\r\n", "\n").replace("\r", "\n")
        if len(value) > XLSX_MAX_CELL_CHARS:
            value = value[:XLSX_MAX_CELL_CHARS - 1] + "…"
    if value is None or value == "":
        return
    cell = ws.cell(row=row, column=col)
    if isinstance(value, str):
        cell.value = value
        cell.data_type = "s"
    else:
        cell.value = value
        if fmt:
            cell.number_format = fmt
    if bold:
        cell.font = Font(bold=True)


def _put_header(ws: Worksheet, row: int, titles: Sequence[str]) -> None:
    for col, title in enumerate(titles, start=1):
        _put(ws, row, col, title, bold=True)


def _table_sheet(ws: Worksheet, columns: Sequence[_Column], items: Sequence[Any],
                 total: Any = None) -> None:
    """Bảng: tiêu đề, dữ liệu, (dòng Tổng); freeze + auto_filter (không gồm Tổng)."""
    _put_header(ws, 1, [c[0] for c in columns])
    for r, item in enumerate(items, start=2):
        for c, (_, get, fmt) in enumerate(columns, start=1):
            _put(ws, r, c, get(item), fmt)
    last = len(items) + 1
    if total is not None:
        for c, (title, get, fmt) in enumerate(columns, start=1):
            value = "Tổng" if c == 1 else ("" if title == "Bot" else get(total))
            _put(ws, last + 1, c, value, fmt)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(columns))}{last}"


def _snapshot_text(commit: CommitInfo | None, timezone: str) -> str:
    if commit is None:
        return "Không có commit trước start (LOC start = 0)"
    local = commit.committer_date.astimezone(ZoneInfo(timezone))
    return f'{commit.short_sha} ({local:%Y-%m-%d %H:%M}) "{commit.subject}"'


def _check_text(check: CrossCheck) -> str:
    return "Đạt" if check.ok else "Lệch quá ngưỡng"


def _overview_sheet(ws: Worksheet, data: ReportData) -> None:
    """Tong_quan: bảng Mục / Giá trị rồi bảng export (mục 5.3 bước 2)."""
    period, check = data.period, data.cross_check
    all_label = "Tất cả user" + (" (loại trừ bot)" if data.options.exclude_bots else "")
    user_text = all_label if data.all_users else ", ".join(sorted(data.users))
    items: list[tuple[str, Any, str | None]] = [
        ("Repo", data.repo, None),
        ("Nhánh", data.branch, None),
        ("Múi giờ", period.timezone, None),
        ("Start", period.start.isoformat(), None),
        ("End", period.end.isoformat(), None),
        ("C_START", data.c_start.sha if data.c_start else "(rỗng)", None),
        ("Snapshot start", _snapshot_text(data.c_start, period.timezone), None),
        ("C_END", data.c_end.sha, None),
        ("Snapshot end", _snapshot_text(data.c_end, period.timezone), None),
        ("User", user_text, None),
        ("Số user được tính", len(data.users), _INT),
        ("LOC start", data.loc_start, _INT),
        ("LOC end", data.loc_end, _INT),
        ("Chênh lệch", data.loc_delta, _DELTA),
        ("Commit", total_metrics(list(data.user_rows)).commits, _INT),
        ("Tổng code ròng", check.sum_net_code, _DELTA),
        ("LOC end kỳ vọng", check.expected_loc_end, _INT),
        ("Lệch kiểm tra chéo (%)", check.diff_percent, _PERCENT),
        ("Ngưỡng (%)", check.threshold_percent, _PERCENT),
        ("Kiểm tra chéo", _check_text(check), None),
        ("Số cảnh báo", len(data.warnings), _INT),
        ("Phiên bản công cụ", data.tool_version, None),
    ]
    if any(e.role != "start" and e.accounts is not None for e in data.exports):
        items.append(("Ghi chú", _NOTE_USER_EXPORT, None))
    _put_header(ws, 1, ["Mục", "Giá trị"])
    for r, (label, value, fmt) in enumerate(items, start=2):
        _put(ws, r, 1, label)
        _put(ws, r, 2, value, fmt)
    row = len(items) + 3
    _put_header(ws, row, ["File export", "Loại", "User", "Số file", "LOC", "LOC end kỳ vọng",
                          "Lệch (%)", "Kiểm tra chéo"])
    roles = {"start": "Start", "end": "End", "separate": "End (riêng)"}
    for exp in data.exports:
        row += 1
        if exp.role == "start":
            who = ""
        elif exp.tag == "all" or exp.accounts is None:
            who = all_label
        else:
            who = ", ".join(exp.accounts)
        values: list[tuple[Any, str | None]] = [
            (exp.file_name, None), (roles[exp.role], None), (who, None),
            (exp.files, _INT), (exp.loc, _INT)]
        if exp.cross_check is not None:
            values += [(exp.cross_check.expected_loc_end, _INT),
                       (exp.cross_check.diff_percent, _PERCENT),
                       (_check_text(exp.cross_check), None)]
        for c, (value, fmt) in enumerate(values, start=1):
            _put(ws, row, c, value, fmt)


def _fit_columns(ws: Worksheet) -> None:
    """Độ rộng cột tất định theo độ dài giá trị (10..60)."""
    widths: dict[str, int] = {}
    for row in ws.iter_rows():
        for cell in row:
            if cell.value is not None:
                letter = cell.column_letter
                widths[letter] = max(widths.get(letter, 0), len(str(cell.value)))
    for letter, width in sorted(widths.items()):
        ws.column_dimensions[letter].width = min(60, max(10, width + 2))


def _normalize_xlsx(raw: bytes, c_end: CommitInfo, timezone: str) -> bytes:
    """Ghi lại zip của openpyxl: giữ thứ tự entry, giờ cố định = C_END, sửa core.xml (N14)."""
    stamp = c_end.committer_date.astimezone(ZoneInfo(timezone)).replace(tzinfo=None)
    date_time = max(stamp, _ZIP_EPOCH).timetuple()[:6]
    core_time = c_end.committer_date.astimezone(UTC).strftime("%Y-%m-%dT%H:%M:%SZ")
    out = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(raw)) as src, zipfile.ZipFile(
            out, "w", compression=zipfile.ZIP_DEFLATED, compresslevel=6) as dst:
        for info in src.infolist():
            content = src.read(info.filename)
            if info.filename == "docProps/core.xml":
                text = _CORE_TIME_RE.sub(lambda m: f"{m.group(1)}{core_time}{m.group(3)}",
                                         content.decode("utf-8"))
                content = text.encode("utf-8")
            entry = zipfile.ZipInfo(info.filename, date_time)
            entry.compress_type = zipfile.ZIP_DEFLATED
            entry.create_system = 3
            entry.external_attr = 0o100644 << 16
            dst.writestr(entry, content, compresslevel=6)
    return out.getvalue()


def write_xlsx(path: Path, data: ReportData) -> Path:
    """loc_report.xlsx 6 sheet, chuẩn hóa tái lập (D40-N14)."""
    wb = Workbook()
    sheets = [wb.active]
    sheets[0].title = SHEET_NAMES[0]
    sheets += [wb.create_sheet(name) for name in SHEET_NAMES[1:]]
    overview, languages, users, files, weeks, warns = sheets
    _overview_sheet(overview, data)
    user_rows = _sorted_users(data.user_rows)
    _table_sheet(languages, _LANGUAGE_COLUMNS, _sorted_languages(data.languages))
    _table_sheet(users, _USER_COLUMNS, user_rows, total=total_metrics(user_rows))
    _table_sheet(files, _FILE_COLUMNS, _sorted_files(data.files))
    _table_sheet(weeks, _WEEK_COLUMNS, _sorted_weeks(data.weeks))
    _table_sheet(warns, _WARNING_COLUMNS, _sorted_warnings(data.warnings))
    for ws in sheets:
        _fit_columns(ws)
    buf = io.BytesIO()
    wb.save(buf)
    return _write_bytes(Path(path), _normalize_xlsx(buf.getvalue(), data.c_end,
                                                    data.period.timezone))


# ---------- CSV ----------
def _csv_cell(value: Any) -> Any:
    """Ô CSV: văn bản làm sạch + tiền tố ' khi bắt đầu bằng ký tự công thức; số ghi thô."""
    if not isinstance(value, str):
        return value
    text = clean_text(value)
    return "'" + text if text.startswith(_CSV_DANGEROUS) else text


def _write_csv(path: Path, columns: Sequence[_Column], items: Iterable[Any]) -> Path:
    """CSV UTF-8 có BOM, dấu phẩy, \\r\\n, tiêu đề tiếng Việt (D40-N20)."""
    buf = io.StringIO()
    writer = csv.writer(buf, delimiter=",", lineterminator="\r\n", quoting=csv.QUOTE_MINIMAL)
    writer.writerow([_csv_cell(c[0]) for c in columns])
    for item in items:
        writer.writerow([_csv_cell(get(item)) for _, get, _ in columns])
    return _write_bytes(Path(path), buf.getvalue().encode("utf-8-sig"))


def write_csv_reports(out_dir: Path, data: ReportData) -> list[Path]:
    """loc_by_user.csv, loc_by_language.csv, loc_by_file.csv (thứ tự này)."""
    out_dir = Path(out_dir)
    return [
        _write_csv(out_dir / USER_CSV_NAME, _USER_COLUMNS, _sorted_users(data.user_rows)),
        _write_csv(out_dir / LANGUAGE_CSV_NAME, _LANGUAGE_COLUMNS,
                   _sorted_languages(data.languages)),
        _write_csv(out_dir / FILE_CSV_NAME, _FILE_COLUMNS, _sorted_files(data.files)),
    ]


def write_warnings_csv(path: Path, warnings: Sequence[ReportWarning]) -> Path:
    """warnings.csv (cột Loại, Đường dẫn, Nội dung, Phạm vi)."""
    return _write_csv(Path(path), _WARNING_COLUMNS, _sorted_warnings(warnings))


# ---------- manifest ----------
def _cross_check_dict(check: CrossCheck) -> dict[str, Any]:
    return {
        "expected_loc_end": check.expected_loc_end,
        "diff_percent": check.diff_percent,
        "loc_end": check.loc_end,
        "sum_net_code": check.sum_net_code,
        "threshold_percent": check.threshold_percent,
        "ok": check.ok,
    }


def _export_dict(exp: ExportSummary) -> dict[str, Any]:
    return {
        "role": exp.role,
        "file": exp.file_name,
        "tag": exp.tag,
        "accounts": list(exp.accounts) if exp.accounts is not None else None,
        "files": exp.files,
        "loc": exp.loc,
        "cross_check": _cross_check_dict(exp.cross_check) if exp.cross_check else None,
    }


def build_manifest(data: ReportData, generated_at: datetime | None = None) -> dict[str, Any]:
    """Dict manifest theo thứ tự key §12 + options, exports (D40-N15)."""
    if generated_at is None:
        generated_at = datetime.now(ZoneInfo(data.period.timezone)).replace(microsecond=0)
    elif generated_at.tzinfo is None or generated_at.utcoffset() is None:
        raise ValueError("generated_at phải có múi giờ (aware datetime)")
    selected = data.options.selected_users
    return {
        "tool_version": data.tool_version,
        "generated_at": generated_at.isoformat(timespec="seconds"),
        "repo": data.repo,
        "branch": data.branch,
        "timezone": data.period.timezone,
        "start": data.period.start.isoformat(),
        "end": data.period.end.isoformat(),
        "c_start": data.c_start.sha if data.c_start else None,
        "c_end": data.c_end.sha,
        "users": sorted(data.users),
        "all_users": data.all_users,
        "config": config_to_dict(data.config),
        "totals": {"loc_start": data.loc_start, "loc_end": data.loc_end,
                   "loc_delta": data.loc_delta},
        "cross_check": _cross_check_dict(data.cross_check),
        "options": {
            "selected_users": list(selected) if selected is not None else None,
            "exclude_bots": data.options.exclude_bots,
            "separate_users": data.options.separate_users,
        },
        "exports": [_export_dict(e) for e in data.exports],
    }


def _json_safe(value: Any) -> Any:
    """Thay surrogate lẻ trong mọi chuỗi (key và giá trị, đệ quy) bằng U+FFFD (A13)."""
    if isinstance(value, str):
        return _SURROGATE_RE.sub("\N{REPLACEMENT CHARACTER}", value)
    if isinstance(value, Mapping):
        return {_json_safe(k): _json_safe(v) for k, v in value.items()}
    if isinstance(value, list | tuple):
        return [_json_safe(v) for v in value]
    return value


def write_manifest(path: Path, manifest: Mapping[str, Any]) -> Path:
    """JSON indent=2, ensure_ascii=False, UTF-8 không BOM, LF, có newline cuối."""
    text = json.dumps(_json_safe(manifest), indent=2, ensure_ascii=False) + "\n"
    return _write_bytes(Path(path), text.encode("utf-8"))


def write_reports(out_dir: Path, data: ReportData,
                  generated_at: datetime | None = None) -> list[Path]:
    """Ghi đủ 6 file vào out_dir; manifest ghi cuối cùng. Trả đường dẫn theo thứ tự cố định."""
    out_dir = Path(out_dir)
    if not out_dir.is_dir():
        logger.error("Thư mục kết quả không tồn tại hoặc không phải thư mục: %s", out_dir)
        raise ReportError(f"Thư mục kết quả không tồn tại hoặc không phải thư mục: {out_dir}")
    manifest = build_manifest(data, generated_at)
    paths = [write_xlsx(out_dir / XLSX_NAME, data)]
    paths += write_csv_reports(out_dir, data)
    paths.append(write_warnings_csv(out_dir / WARNINGS_CSV_NAME, data.warnings))
    paths.append(write_manifest(out_dir / MANIFEST_NAME, manifest))
    logger.info("Đã ghi báo cáo vào %s", out_dir)
    return paths
